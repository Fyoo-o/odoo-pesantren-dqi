from odoo import api, fields, models, _
from odoo.exceptions import UserError
import base64
import csv
import io
from datetime import date
import calendar

try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, Inches, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


BULAN_SELECTION = [
    ('1', 'Januari'), ('2', 'Februari'), ('3', 'Maret'),
    ('4', 'April'), ('5', 'Mei'), ('6', 'Juni'),
    ('7', 'Juli'), ('8', 'Agustus'), ('9', 'September'),
    ('10', 'Oktober'), ('11', 'November'), ('12', 'Desember'),
]

BULAN_DICT = dict(BULAN_SELECTION)


class WizardRekapTahsinTahfidz(models.TransientModel):
    _name = 'cdn.wizard_rekap_tahsin_tahfidz'
    _description = 'Wizard Rekap Tahsin Tahfidz (Laporan Pencapaian Per Halaqoh)'

    def _domain_halaqoh_id(self):
        tahun_ajaran = self.env.user.company_id.tahun_ajaran_aktif.id
        user = self.env.user

        if user.has_group('pesantren_kesantrian.group_kesantrian_manager'):
            return [('fiscalyear_id', '=', tahun_ajaran)]

        employee = self.env['hr.employee'].search(
            [('user_id', '=', user.id)], limit=1)
        if employee:
            return [
                ('fiscalyear_id', '=', tahun_ajaran),
                '|',
                ('penanggung_jawab_id', '=', employee.id),
                ('pengganti_ids', 'in', [employee.id])
            ]
        return [('id', '=', 0)]

    halaqoh_id = fields.Many2one(
        'cdn.halaqoh', string='Halaqoh', required=True,
        domain=_domain_halaqoh_id)
    bulan = fields.Selection(
        BULAN_SELECTION, string='Bulan', required=True,
        default=lambda self: str(date.today().month))
    tahun = fields.Char(
        string='Tahun', required=True,
        default=lambda self: str(date.today().year))

    # Auto-filled from halaqoh
    pembimbing_id = fields.Many2one(
        'hr.employee', string='Pembimbing Halaqah',
        related='halaqoh_id.penanggung_jawab_id', readonly=True)

    # Inline results
    rekap_line_ids = fields.One2many(
        'cdn.wizard_rekap_tahsin_tahfidz_line', 'wizard_id',
        string='Detail Pencapaian', readonly=True)

    # Export
    data_file = fields.Binary(string='File')
    file_name = fields.Char(string='Nama File')

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        bulan = res.get('bulan', str(date.today().month))
        tahun = res.get('tahun', str(date.today().year))

        halaqoh_id = res.get('halaqoh_id')
        if not halaqoh_id:
            tahun_ajaran = self.env.user.company_id.tahun_ajaran_aktif.id
            employee = self.env['hr.employee'].search([('user_id', '=', self.env.uid)], limit=1)
            domain = [('fiscalyear_id', '=', tahun_ajaran)]
            if employee and not self.env.user.has_group('pesantren_kesantrian.group_kesantrian_manager'):
                domain.extend(['|', ('penanggung_jawab_id', '=', employee.id), ('pengganti_ids', 'in', [employee.id])])
            halaqoh = self.env['cdn.halaqoh'].search(domain, limit=1)
            if halaqoh:
                halaqoh_id = halaqoh.id
                res['halaqoh_id'] = halaqoh_id

        if halaqoh_id and bulan and tahun:
            lines = self._build_rekap_lines(halaqoh_id, int(bulan), int(tahun))
            res['rekap_line_ids'] = lines

        return res

    @api.onchange('halaqoh_id', 'bulan', 'tahun')
    def _onchange_rekap_params(self):
        if not (self.halaqoh_id and self.bulan and self.tahun):
            self.rekap_line_ids = [(5, 0, 0)]
            return

        lines = self._build_rekap_lines(self.halaqoh_id.id, int(self.bulan), int(self.tahun))
        self.rekap_line_ids = [(5, 0, 0)] + lines

    def _build_rekap_lines(self, halaqoh_id, bulan_int, tahun_int):
        _, last_day = calendar.monthrange(tahun_int, bulan_int)
        tgl_awal = date(tahun_int, bulan_int, 1)
        tgl_akhir = date(tahun_int, bulan_int, last_day)

        halaqoh = self.env['cdn.halaqoh'].browse(halaqoh_id)
        lines = []
        no = 1

        for siswa in halaqoh.siswa_ids.sorted(key=lambda s: s.name or ''):
            penilaian_records = self.env['cdn.penilaian_quran'].search([
                ('siswa_id', '=', siswa.id),
                ('tanggal', '>=', tgl_awal),
                ('tanggal', '<=', tgl_akhir),
                ('state', '=', 'done')
            ], order='tanggal asc')

            tahsin_parts = []
            for rec in penilaian_records:
                if rec.buku_harian_id:
                    parts = []
                    if rec.buku_harian_id:
                        parts.append(rec.buku_harian_id.name)
                    if rec.jilid_harian_id:
                        parts.append(rec.jilid_harian_id.name)
                    if rec.halaman_harian:
                        parts.append(f"Hal {rec.halaman_harian}")
                    tahsin_parts.append(' '.join(parts))
            tahsin_str = tahsin_parts[-1] if tahsin_parts else '-'

            murojaah_parts = []
            for rec in penilaian_records:
                if rec.juz_murajaah:
                    parts = []
                    parts.append(f"Juz {rec.juz_murajaah}")
                    if rec.surah_murajaah_id:
                        parts.append(rec.surah_murajaah_id.name)
                    if rec.halaman_murajaah:
                        parts.append(f"Hal {rec.halaman_murajaah}")
                    murojaah_parts.append(' '.join(parts))
            murojaah_str = murojaah_parts[-1] if murojaah_parts else '-'

            total_maqra_bulan = 0
            for rec in penilaian_records:
                for line in rec.tahfidz_line_ids:
                    total_maqra_bulan += line.jml_baris or 0
            hafalan_baru_str = f"{total_maqra_bulan} Maqra" if total_maqra_bulan > 0 else '-'

            all_penilaian = self.env['cdn.penilaian_quran'].search([
                ('siswa_id', '=', siswa.id),
                ('state', '=', 'done')
            ])
            total_maqra_all = 0
            for rec in all_penilaian:
                for line in rec.tahfidz_line_ids:
                    total_maqra_all += line.jml_baris or 0
            jml_semua_str = f"{total_maqra_all} Maqra" if total_maqra_all > 0 else '-'

            nilai_list = []
            for rec in penilaian_records:
                for line in rec.tahfidz_line_ids:
                    if line.nilai_hafalan:
                        nilai_list.append(line.nilai_hafalan)
            nilai_avg = round(sum(nilai_list) / len(nilai_list)) if nilai_list else 0

            keterangan_parts = []
            for rec in penilaian_records:
                if rec.keterangan_tahfidz:
                    keterangan_parts.append(rec.keterangan_tahfidz)
            keterangan_str = '; '.join(keterangan_parts) if keterangan_parts else ''

            lines.append((0, 0, {
                'no': no,
                'siswa_id': siswa.id,
                'nama': siswa.name,
                'tahsin': tahsin_str,
                'murojaah': murojaah_str,
                'hafalan_baru': hafalan_baru_str,
                'jml_semua_hafalan': jml_semua_str,
                'nilai': nilai_avg,
                'keterangan': keterangan_str,
            }))
            no += 1

        return lines

    def action_export_csv(self):
        self._onchange_rekap_params()
        if not self.rekap_line_ids:
            raise UserError(_("Tidak ada data untuk diekspor."))

        output = io.StringIO()
        writer = csv.writer(output, delimiter=',',
                            quotechar='"', quoting=csv.QUOTE_MINIMAL)

        writer.writerow(['LAPORAN PENCAPAIAN TAHFIZH SANTRI/WATI'])
        writer.writerow([self.env.user.company_id.name or ''])
        writer.writerow([f'Pembimbing Halaqah: {self.pembimbing_id.name or "-"}'])
        writer.writerow([f'Halaqoh: {self.halaqoh_id.name}'])
        writer.writerow([f'Bulan: {BULAN_DICT.get(self.bulan, "")}'])
        writer.writerow([f'Tahun: {self.tahun}'])
        writer.writerow([])

        writer.writerow(['NO', 'NAMA', 'TAHSIN', 'MUROJAAH',
                         'HAFALAN BARU', 'JUMLAH SEMUA HAFALAN',
                         'NILAI', 'KETERANGAN'])

        for line in self.rekap_line_ids:
            writer.writerow([
                line.no, line.nama, line.tahsin, line.murojaah,
                line.hafalan_baru, line.jml_semua_hafalan,
                line.nilai, line.keterangan or ''
            ])

        csv_data = output.getvalue().encode('utf-8')
        nama_bulan = BULAN_DICT.get(self.bulan, '')
        file_name = f"Rekap_TahsinTahfidz_{self.halaqoh_id.name}_{nama_bulan}_{self.tahun}.csv"

        self.write({
            'data_file': base64.b64encode(csv_data),
            'file_name': file_name
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/?model={self._name}&id={self.id}&field=data_file&download=true&filename={self.file_name}',
            'target': 'self',
        }

    def action_export_docx(self):
        if not HAS_DOCX:
            raise UserError(
                _("Modul 'python-docx' tidak ditemukan. Silakan hubungi administrator."))

        self._onchange_rekap_params()
        if not self.rekap_line_ids:
            raise UserError(_("Tidak ada data untuk diekspor."))

        doc = DocxDocument()

        # Page margins
        for section in doc.sections:
            section.top_margin = Cm(1.5)
            section.bottom_margin = Cm(1.5)
            section.left_margin = Cm(1.5)
            section.right_margin = Cm(1.5)

        nama_bulan = BULAN_DICT.get(self.bulan, '')
        company_name = self.env.user.company_id.name or ''

        # ---- Title ----
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(0)
        r_title = p_title.add_run('LAPORAN PENCAPAIAN TAHFIZH SANTRI/WATI')
        r_title.bold = True
        r_title.font.size = Pt(14)
        r_title.font.name = 'Times New Roman'

        p_company = doc.add_paragraph()
        p_company.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_company.paragraph_format.space_before = Pt(0)
        p_company.paragraph_format.space_after = Pt(12)
        r_company = p_company.add_run(company_name.upper())
        r_company.bold = True
        r_company.font.size = Pt(12)
        r_company.font.name = 'Times New Roman'

        # ---- Info Lines ----
        info_data = [
            ('PEMBIMBING HALAQAH', self.pembimbing_id.name or '-'),
            ('HP / WA', '-'),
            ('KELAS', 'TAHSIN-TAHFIDZ'),
            ('BULAN', nama_bulan.upper()),
            ('TAHUN', str(self.tahun)),
        ]
        for label, value in info_data:
            p_info = doc.add_paragraph()
            p_info.paragraph_format.space_before = Pt(0)
            p_info.paragraph_format.space_after = Pt(2)
            p_info.paragraph_format.line_spacing = 1.0
            r_label = p_info.add_run(f'{label:<24}')
            r_label.bold = True
            r_label.font.size = Pt(10)
            r_label.font.name = 'Times New Roman'
            r_val = p_info.add_run(f': {value}')
            r_val.font.size = Pt(10)
            r_val.font.name = 'Times New Roman'

        # ---- Subtitle ----
        doc.add_paragraph()
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(6)
        r_sub = p_sub.add_run('PROGRAM TAHSIN TAHFIZH')
        r_sub.bold = True
        r_sub.font.size = Pt(11)
        r_sub.font.name = 'Times New Roman'

        # ---- Data Table ----
        headers = ['NO', 'NAMA', 'TAHSIN', 'MUROJAAH',
                   'HAFALAN BARU', 'JUMLAH SEMUA HAFALAN',
                   'NILAI', 'KETERANGAN']
        col_widths = [Cm(1.2), Cm(4.0), Cm(2.5), Cm(2.5),
                      Cm(3.0), Cm(3.5), Cm(1.5), Cm(3.0)]

        num_rows = len(self.rekap_line_ids) + 1
        table = doc.add_table(rows=num_rows, cols=8)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Set column widths
        for i, width in enumerate(col_widths):
            for row_obj in table.rows:
                row_obj.cells[i].width = width

        # Header row
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(header)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'
            # Blue background
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="4472C4"/>')
            cell._tc.get_or_add_tcPr().append(shd)

        # Data rows
        for idx, line in enumerate(self.rekap_line_ids):
            row_idx = idx + 1
            data = [
                str(line.no), line.nama or '', line.tahsin or '',
                line.murojaah or '', line.hafalan_baru or '',
                line.jml_semua_hafalan or '', str(line.nilai) if line.nilai else '',
                line.keterangan or ''
            ]
            for col_idx, val in enumerate(data):
                cell = table.rows[row_idx].cells[col_idx]
                cell.text = ''
                p = cell.paragraphs[0]
                if col_idx in [0, 6]:  # NO & NILAI centered
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(val)
                run.font.size = Pt(9)
                run.font.name = 'Times New Roman'

        # ---- Signatures ----
        doc.add_paragraph()
        doc.add_paragraph()
        today = date.today()
        tgl_str = f"{today.day} {BULAN_DICT.get(str(today.month), '')} {today.year}"

        sig_table = doc.add_table(rows=4, cols=2)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_table.autofit = True

        # Remove borders from signature table
        for row_obj in sig_table.rows:
            for cell in row_obj.cells:
                for border_name in ['top', 'bottom', 'left', 'right']:
                    tag = f'w:{border_name}'
                    element = parse_xml(
                        f'<w:tcBorders {nsdecls("w")}>'
                        f'<{tag} w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                        f'</w:tcBorders>'
                    )
                    cell._tc.get_or_add_tcPr().append(element)

        # Row 0: Mengetahui / Tanggal
        for i, text in enumerate(['Mengetahui,', f'Pelaihari, {tgl_str}']):
            p = sig_table.rows[0].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'

        # Row 1: Jabatan
        for i, text in enumerate([f'Kepala {company_name}', 'Pembimbing']):
            p = sig_table.rows[1].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'

        # Row 2: spacer (empty)
        for i in range(2):
            sig_table.rows[2].cells[i].text = ''
            # Add space
            p = sig_table.rows[2].cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(40)

        # Row 3: Names
        pembimbing_name = self.pembimbing_id.name or '______________________________'
        for i, text in enumerate(['(______________________________)', f'({pembimbing_name})']):
            p = sig_table.rows[3].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'
            run.font.underline = True

        # ---- Save & Download ----
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        docx_data = output.read()

        file_name = f"Rekap_TahsinTahfidz_{self.halaqoh_id.name}_{nama_bulan}_{self.tahun}.docx"

        self.write({
            'data_file': base64.b64encode(docx_data),
            'file_name': file_name
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/?model={self._name}&id={self.id}&field=data_file&download=true&filename={self.file_name}',
            'target': 'self',
        }


class WizardRekapTahsinTahfidzLine(models.TransientModel):
    _name = 'cdn.wizard_rekap_tahsin_tahfidz_line'
    _description = 'Line Rekap Tahsin Tahfidz'

    wizard_id = fields.Many2one(
        'cdn.wizard_rekap_tahsin_tahfidz', string='Wizard', ondelete='cascade')
    no = fields.Integer(string='No')
    siswa_id = fields.Many2one('cdn.siswa', string='Santri')
    nama = fields.Char(string='Nama')
    tahsin = fields.Char(string='Tahsin')
    murojaah = fields.Char(string='Murojaah')
    hafalan_baru = fields.Char(string='Hafalan Baru')
    jml_semua_hafalan = fields.Char(string='Jumlah Semua Hafalan')
    nilai = fields.Integer(string='Nilai')
    keterangan = fields.Char(string='Keterangan')
