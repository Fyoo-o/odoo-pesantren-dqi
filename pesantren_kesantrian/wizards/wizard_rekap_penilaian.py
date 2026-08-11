from odoo import api, fields, models, _
from odoo.exceptions import UserError
import base64
import csv
import io

try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class WizardRekapPenilaian(models.TransientModel):
    _name = 'cdn.wizard_rekap_penilaian'
    _description = 'Wizard Rekap Penilaian Santri'

    siswa_id = fields.Many2one('cdn.siswa', string='Nama Siswa', required=True)
    barcode = fields.Char(string='Kartu Santri')
    tgl_awal = fields.Date(string='Tanggal Awal',
                           required=True, default=fields.Date.context_today)
    tgl_akhir = fields.Date(string='Tanggal Akhir',
                            required=True, default=fields.Date.context_today)

    # Inline results
    rekap_line_ids = fields.One2many(
        'cdn.wizard_rekap_penilaian_line', 'wizard_id', string='Detail Penilaian', readonly=True)

    # Export fields
    data_file = fields.Binary(string='File')
    file_name = fields.Char(string='Nama File')

    @api.onchange('barcode')
    def _onchange_barcode(self):
        if self.barcode:
            siswa = self.env['cdn.siswa'].search(
                [('barcode_santri', '=', self.barcode)], limit=1)
            if siswa:
                self.siswa_id = siswa.id
            else:
                self.siswa_id = False
                barcode_sementara = self.barcode
                self.barcode = False
                return {
                    'warning': {
                        'title': "Perhatian !",
                        'message': f"Data Santri dengan Kartu Santri {barcode_sementara} tidak ditemukan."
                    }
                }
        else:
            self.siswa_id = False

    @api.onchange('siswa_id', 'tgl_awal', 'tgl_akhir')
    def _onchange_rekap_params(self):
        if self.siswa_id:
            self.barcode = self.siswa_id.barcode_santri
        else:
            self.barcode = False

        if not (self.siswa_id and self.tgl_awal and self.tgl_akhir):
            self.rekap_line_ids = [(5, 0, 0)]
            return

        if self.tgl_awal > self.tgl_akhir:
            return

        lines = []

        # Penilaian Quran records
        penilaian_records = self.env['cdn.penilaian_quran'].search([
            ('siswa_id', '=', self.siswa_id.id),
            ('tanggal', '>=', self.tgl_awal),
            ('tanggal', '<=', self.tgl_akhir),
            ('state', '=', 'done')
        ], order='tanggal asc')

        for rec in penilaian_records:
            # 1. Tahfidz (Lines)
            for line in rec.tahfidz_line_ids:
                materi = f"{line.surah_id.name} {line.ayat_awal.name}-{line.ayat_akhir.name}"
                lines.append((0, 0, {
                    'tanggal': rec.tanggal,
                    'kegiatan': 'Tahfidz',
                    'materi': materi,
                    'nilai': str(line.nilai_hafalan),
                    'predikat': line.predikat or '-',
                    'keterangan': line.keterangan or '-'
                }))

            # 2. Tahsin
            if rec.buku_harian_id:
                materi_tahsin = f"{rec.buku_harian_id.name} - {rec.jilid_harian_id.name} Hal {rec.halaman_harian}"
                lines.append((0, 0, {
                    'tanggal': rec.tanggal,
                    'kegiatan': 'Tahsin',
                    'materi': materi_tahsin,
                    'nilai': str(rec.nilai_tahsin_harian),
                    'predikat': rec.predikat_tahsin_harian or '-',
                    'keterangan': rec.catatan_harian or '-'
                }))

            # 3. Murajaah
            if rec.juz_murajaah:
                materi_murajaah = f"Juz {rec.juz_murajaah} {rec.surah_murajaah_id.name} Hal {rec.halaman_murajaah}"
                lines.append((0, 0, {
                    'tanggal': rec.tanggal,
                    'kegiatan': 'Murajaah',
                    'materi': materi_murajaah,
                    'nilai': '-',
                    'predikat': '-',
                    'keterangan': rec.catatan_murajaah_harian or '-'
                }))

        self.rekap_line_ids = [(5, 0, 0)] + lines

    def action_export_csv(self):
        self._onchange_rekap_params()
        if not self.rekap_line_ids:
            raise UserError(_("Tidak ada data untuk diekspor."))

        output = io.StringIO()
        writer = csv.writer(output, delimiter=',',
                            quotechar='"', quoting=csv.QUOTE_MINIMAL)

        # Header
        writer.writerow(['No', 'Tanggal', 'Kegiatan', 'Materi',
                        'Nilai', 'Predikat', 'Keterangan'])

        # Data
        for i, line in enumerate(self.rekap_line_ids, 1):
            writer.writerow([
                i,
                line.tanggal,
                line.kegiatan,
                line.materi,
                line.nilai,
                line.predikat,
                line.keterangan
            ])

        csv_data = output.getvalue().encode('utf-8')
        file_name = f"Rekap_Penilaian_{self.siswa_id.name}_{self.tgl_awal}_sd_{self.tgl_akhir}.csv"

        self.write({
            'data_file': base64.b64encode(csv_data),
            'file_name': file_name
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/?model={self._name}&id={self.id}&field=data_file&download=true&filename={self.file_name}',
            'target': 'self',
        }

    def action_export_docx(self):
        if not HAS_DOCX:
            raise UserError(
                _("Modul 'python-docx' tidak ditemukan. Silakan hubungi administrator."))

        self._onchange_rekap_params()
        if not self.rekap_line_ids:
            raise UserError(_("Tidak ada data untuk diekspor."))

        doc = DocxDocument()

        # Page margins
        for section in doc.sections:
            section.top_margin = Cm(1.5)
            section.bottom_margin = Cm(1.5)
            section.left_margin = Cm(1.5)
            section.right_margin = Cm(1.5)

        # ---- Title ----
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(12)
        r_title = p_title.add_run('Rekap Penilaian Santri (Halaqoh)')
        r_title.bold = True
        r_title.font.size = Pt(14)
        r_title.font.name = 'Times New Roman'

        # ---- Info ----
        info_data = [
            ('Nama Siswa', self.siswa_id.name or '-'),
            ('NIS', self.siswa_id.nis or '-'),
            ('Periode', f'{self.tgl_awal} s/d {self.tgl_akhir}'),
        ]
        for label, value in info_data:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            r_l = p.add_run(f'{label:<20}')
            r_l.bold = True
            r_l.font.size = Pt(10)
            r_l.font.name = 'Times New Roman'
            r_v = p.add_run(f': {value}')
            r_v.font.size = Pt(10)
            r_v.font.name = 'Times New Roman'

        doc.add_paragraph()

        # ---- Data Table ----
        headers = ['No', 'Tanggal', 'Kegiatan', 'Materi', 'Nilai', 'Predikat', 'Keterangan']
        col_widths = [Cm(1.0), Cm(2.5), Cm(2.5), Cm(5.0), Cm(1.5), Cm(2.0), Cm(3.5)]

        num_rows = len(self.rekap_line_ids) + 1
        table = doc.add_table(rows=num_rows, cols=7)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Set column widths
        for i, width in enumerate(col_widths):
            for row_obj in table.rows:
                row_obj.cells[i].width = width

        # Header row
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(header)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="D3D3D3"/>')
            cell._tc.get_or_add_tcPr().append(shd)

        # Data rows
        for idx, line in enumerate(self.rekap_line_ids):
            row_idx = idx + 1
            tgl_str = line.tanggal.strftime('%d/%m/%Y') if line.tanggal else '-'
            data = [
                str(idx + 1), tgl_str, line.kegiatan or '',
                line.materi or '', line.nilai or '',
                line.predikat or '', line.keterangan or ''
            ]
            for col_idx, val in enumerate(data):
                cell = table.rows[row_idx].cells[col_idx]
                cell.text = ''
                p = cell.paragraphs[0]
                if col_idx in [0, 4, 5]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(val)
                run.font.size = Pt(9)
                run.font.name = 'Times New Roman'

        # ---- Save & Download ----
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        docx_data = output.read()

        file_name = f"Rekap_Penilaian_{self.siswa_id.name}_{self.tgl_awal}_sd_{self.tgl_akhir}.docx"

        self.write({
            'data_file': base64.b64encode(docx_data),
            'file_name': file_name
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/?model={self._name}&id={self.id}&field=data_file&download=true&filename={self.file_name}',
            'target': 'self',
        }


class WizardRekapPenilaianLine(models.TransientModel):
    _name = 'cdn.wizard_rekap_penilaian_line'
    _description = 'Line Rekap Penilaian Santri'

    wizard_id = fields.Many2one(
        'cdn.wizard_rekap_penilaian', string='Wizard', ondelete='cascade')
    tanggal = fields.Date(string='Tanggal')
    kegiatan = fields.Char(string='Kegiatan')
    materi = fields.Char(string='Materi')
    nilai = fields.Char(string='Nilai')
    predikat = fields.Char(string='Predikat')
    keterangan = fields.Char(string='Keterangan')
