from odoo import api, fields, models, _
from odoo.exceptions import UserError
import base64
import csv
import io
from datetime import date

try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class WizardRekapAbsensi(models.TransientModel):
    _name = 'cdn.wizard_rekap_absensi'
    _description = 'Wizard Rekap Absensi Santri'

    siswa_id = fields.Many2one('cdn.siswa', string='Nama Siswa', required=True)
    barcode = fields.Char(string='Kartu Santri')
    tgl_awal = fields.Date(string='Tanggal Awal',
                           required=True, default=fields.Date.context_today)
    tgl_akhir = fields.Date(string='Tanggal Akhir',
                            required=True, default=fields.Date.context_today)

    # Inline results
    rekap_line_ids = fields.One2many(
        'cdn.wizard_rekap_absensi_line', 'wizard_id', string='Detail Kehadiran', readonly=True)

    # Summary fields
    jml_hadir = fields.Integer(string='Hadir', compute='_compute_summary')
    jml_izin = fields.Integer(string='Izin', compute='_compute_summary')
    jml_sakit = fields.Integer(string='Sakit', compute='_compute_summary')
    jml_alpa = fields.Integer(string='Alpa', compute='_compute_summary')
    jml_pulang_sakit = fields.Integer(string='Pulang Sakit', compute='_compute_summary')
    jml_pulang_izin = fields.Integer(string='Pulang Izin', compute='_compute_summary')
    jml_pulang_alpa = fields.Integer(string='Pulang Alpa', compute='_compute_summary')
    jml_guru_pengganti = fields.Integer(string='Guru Pengganti', compute='_compute_summary')
    jml_keluar = fields.Integer(
        string='Izin Keluar', compute='_compute_summary')

    # Export fields
    data_file = fields.Binary(string='File')
    file_name = fields.Char(string='Nama File')

    @api.depends('rekap_line_ids')
    def _compute_summary(self):
        for rec in self:
            rec.jml_hadir = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'Hadir'))
            rec.jml_izin = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'Izin'))
            rec.jml_sakit = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'Sakit'))
            rec.jml_alpa = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'Alpa'))
            rec.jml_pulang_sakit = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'Pulang-Sakit'))
            rec.jml_pulang_izin = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'Pulang-Izin'))
            rec.jml_pulang_alpa = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'Pulang-Alpa'))
            rec.jml_guru_pengganti = len(rec.rekap_line_ids.filtered(
                lambda x: 'Guru Pengganti' in (x.pengabsen or '')))
            rec.jml_keluar = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'keluar'))

    @api.onchange('barcode')
    def _onchange_barcode(self):
        if self.barcode:
            siswa = self.env['cdn.siswa'].search(
                [('barcode_santri', '=', self.barcode)], limit=1)
            if siswa:
                self.siswa_id = siswa.id
            else:
                self.siswa_id = False
                barcode_sementara = self.barcode
                self.barcode = False
                return {
                    'warning': {
                        'title': "Perhatian !",
                        'message': f"Data Santri dengan Kartu Santri {barcode_sementara} tidak ditemukan."
                    }
                }
        else:
            self.siswa_id = False

    @api.onchange('siswa_id', 'tgl_awal', 'tgl_akhir')
    def _onchange_rekap_params(self):
        if self.siswa_id:
            self.barcode = self.siswa_id.barcode_santri
        else:
            self.barcode = False

        if not (self.siswa_id and self.tgl_awal and self.tgl_akhir):
            self.rekap_line_ids = [(5, 0, 0)]
            return

        if self.tgl_awal > self.tgl_akhir:
            return

        lines = []

        # 1. Halaqoh
        halaqoh_lines = self.env['cdn.absen_halaqoh_line'].sudo().search([
            ('siswa_id', '=', self.siswa_id.id),
            ('tanggal', '>=', self.tgl_awal),
            ('tanggal', '<=', self.tgl_akhir)
        ], order='tanggal asc')
        for line in halaqoh_lines:
            absen_hdr = line.absen_id if line.absen_id else False
            is_pengganti = getattr(absen_hdr, 'is_guru_pengganti', False) if absen_hdr else False
            ustadz = getattr(absen_hdr, 'ustadz_id', False)
            pj = getattr(absen_hdr, 'penanggung_jawab_id', False)
            
            pengabsen_str = ustadz.name if ustadz else '-'
            if ustadz and (is_pengganti or (pj and ustadz.id != pj.id)):
                pengabsen_str += ' (Guru Pengganti)'
            lines.append((0, 0, {
                'tanggal': line.tanggal,
                'jenis': 'Halaqoh',
                'kehadiran': line.kehadiran,
                'pengabsen': pengabsen_str,
                'keterangan': line.keterangan or '-'
            }))

        # 2. Malam
        malam_lines = self.env['cdn.absensi_malam_line'].sudo().search([
            ('siswa_id', '=', self.siswa_id.id),
            ('tanggal', '>=', self.tgl_awal),
            ('tanggal', '<=', self.tgl_akhir)
        ], order='tanggal asc')
        for line in malam_lines:
            lines.append((0, 0, {
                'tanggal': line.tanggal,
                'jenis': 'Malam',
                'kehadiran': line.kehadiran_absen,
                'pengabsen': '-',
                'keterangan': line.keterangan or '-'
            }))

        # 3. Tahfidz
        tahfidz_lines = self.env['cdn.absen_tahfidz_quran_line'].sudo().search([
            ('siswa_id', '=', self.siswa_id.id),
            ('tanggal', '>=', self.tgl_awal),
            ('tanggal', '<=', self.tgl_akhir)
        ], order='tanggal asc')
        for line in tahfidz_lines:
            lines.append((0, 0, {
                'tanggal': line.tanggal,
                'jenis': 'Tahfidz',
                'kehadiran': line.kehadiran,
                'pengabsen': '-',
                'keterangan': line.keterangan or '-'
            }))

        # 4. Tahsin
        tahsin_lines = self.env['cdn.absen_tahsin_quran_line'].sudo().search([
            ('siswa_id', '=', self.siswa_id.id),
            ('tanggal', '>=', self.tgl_awal),
            ('tanggal', '<=', self.tgl_akhir)
        ], order='tanggal asc')
        for line in tahsin_lines:
            lines.append((0, 0, {
                'tanggal': line.tanggal,
                'jenis': 'Tahsin',
                'kehadiran': line.kehadiran,
                'pengabsen': '-',
                'keterangan': line.keterangan or '-'
            }))

        # Sort lines by date
        lines.sort(key=lambda x: x[2]['tanggal'])

        self.rekap_line_ids = [(5, 0, 0)] + lines

    def action_export_csv(self):
        self._onchange_rekap_params()  # Ensure lines are populated even if called directly
        if not self.rekap_line_ids:
            raise UserError(_("Tidak ada data untuk diekspor."))

        output = io.StringIO()
        writer = csv.writer(output, delimiter=',',
                            quotechar='"', quoting=csv.QUOTE_MINIMAL)

        # Header
        writer.writerow(['No', 'Tanggal', 'Kegiatan',
                        'Kehadiran', 'Guru / Pengabsen', 'Keterangan'])

        # Data
        for i, line in enumerate(self.rekap_line_ids, 1):
            writer.writerow([
                i,
                line.tanggal,
                line.jenis,
                dict(line._fields['kehadiran'].selection).get(
                    line.kehadiran, line.kehadiran),
                line.pengabsen or '-',
                line.keterangan or '-'
            ])

        csv_data = output.getvalue().encode('utf-8')
        file_name = f"Rekap_Absensi_{self.siswa_id.name}_{self.tgl_awal}_sd_{self.tgl_akhir}.csv"

        self.write({
            'data_file': base64.b64encode(csv_data),
            'file_name': file_name
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/?model={self._name}&id={self.id}&field=data_file&download=true&filename={self.file_name}',
            'target': 'self',
        }

    def action_export_docx(self):
        if not HAS_DOCX:
            raise UserError(
                _("Modul 'python-docx' tidak ditemukan. Silakan hubungi administrator."))

        self._onchange_rekap_params()
        if not self.rekap_line_ids:
            raise UserError(_("Tidak ada data untuk diekspor."))

        doc = DocxDocument()

        # Page margins
        for section in doc.sections:
            section.top_margin = Cm(1.5)
            section.bottom_margin = Cm(1.5)
            section.left_margin = Cm(1.5)
            section.right_margin = Cm(1.5)

        company_name = self.env.user.company_id.name or ''

        # ---- Title ----
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(0)
        r_title = p_title.add_run('REKAP ABSENSI SANTRI')
        r_title.bold = True
        r_title.font.size = Pt(14)
        r_title.font.name = 'Times New Roman'

        p_company = doc.add_paragraph()
        p_company.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_company.paragraph_format.space_before = Pt(0)
        p_company.paragraph_format.space_after = Pt(12)
        r_company = p_company.add_run(company_name.upper())
        r_company.bold = True
        r_company.font.size = Pt(12)
        r_company.font.name = 'Times New Roman'

        # ---- Info Lines ----
        info_data = [
            ('NAMA SISWA', self.siswa_id.name or '-'),
            ('NIS', self.siswa_id.nis or '-'),
            ('PERIODE', f'{self.tgl_awal} s/d {self.tgl_akhir}'),
        ]
        for label, value in info_data:
            p_info = doc.add_paragraph()
            p_info.paragraph_format.space_before = Pt(0)
            p_info.paragraph_format.space_after = Pt(2)
            p_info.paragraph_format.line_spacing = 1.0
            r_label = p_info.add_run(f'{label:<24}')
            r_label.bold = True
            r_label.font.size = Pt(10)
            r_label.font.name = 'Times New Roman'
            r_val = p_info.add_run(f': {value}')
            r_val.font.size = Pt(10)
            r_val.font.name = 'Times New Roman'

        doc.add_paragraph()

        # ---- Data Table ----
        headers = ['NO', 'TANGGAL', 'KEGIATAN', 'KEHADIRAN', 'GURU / PENGABSEN', 'KETERANGAN']
        col_widths = [Cm(1.2), Cm(3.0), Cm(2.5), Cm(2.5), Cm(4.5), Cm(4.0)]

        num_rows = len(self.rekap_line_ids) + 1
        table = doc.add_table(rows=num_rows, cols=6)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Set column widths
        for i, width in enumerate(col_widths):
            for row_obj in table.rows:
                row_obj.cells[i].width = width

        # Header row
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(header)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'
            # Blue background
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="4472C4"/>')
            cell._tc.get_or_add_tcPr().append(shd)

        # Data rows
        for idx, line in enumerate(self.rekap_line_ids):
            row_idx = idx + 1
            keh_label = dict(line._fields['kehadiran'].selection).get(
                line.kehadiran, line.kehadiran) if line.kehadiran else '-'
            data = [
                str(idx + 1),
                str(line.tanggal) if line.tanggal else '',
                line.jenis or '',
                keh_label,
                line.pengabsen or '-',
                line.keterangan or '-'
            ]
            for col_idx, val in enumerate(data):
                cell = table.rows[row_idx].cells[col_idx]
                cell.text = ''
                p = cell.paragraphs[0]
                if col_idx in [0, 1, 3]:  # NO, TANGGAL, KEHADIRAN centered
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(val)
                run.font.size = Pt(9)
                run.font.name = 'Times New Roman'

        # ---- Summary Table ----
        doc.add_paragraph()
        p_sum = doc.add_paragraph()
        p_sum.paragraph_format.space_after = Pt(6)
        r_sum = p_sum.add_run('RINGKASAN KEHADIRAN')
        r_sum.bold = True
        r_sum.font.size = Pt(11)
        r_sum.font.name = 'Times New Roman'

        summary_data = [
            ('Hadir', self.jml_hadir),
            ('Izin', self.jml_izin),
            ('Sakit', self.jml_sakit),
            ('Alpa', self.jml_alpa),
            ('Pulang Sakit', self.jml_pulang_sakit),
            ('Pulang Izin', self.jml_pulang_izin),
            ('Pulang Alpa', self.jml_pulang_alpa),
            ('Sesi Guru Pengganti', self.jml_guru_pengganti),
            ('Izin Keluar', self.jml_keluar),
        ]

        sum_table = doc.add_table(rows=len(summary_data) + 1, cols=2)
        sum_table.style = 'Table Grid'
        sum_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        sum_table.autofit = False

        # Summary column widths
        for row_obj in sum_table.rows:
            row_obj.cells[0].width = Cm(5.0)
            row_obj.cells[1].width = Cm(2.5)

        # Summary header
        for i, header in enumerate(['STATUS KEHADIRAN', 'JUMLAH']):
            cell = sum_table.rows[0].cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(header)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="4472C4"/>')
            cell._tc.get_or_add_tcPr().append(shd)

        # Summary data rows
        for idx, (label, val) in enumerate(summary_data):
            row_idx = idx + 1
            cell_label = sum_table.rows[row_idx].cells[0]
            cell_label.text = ''
            p = cell_label.paragraphs[0]
            run = p.add_run(label)
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'

            cell_val = sum_table.rows[row_idx].cells[1]
            cell_val.text = ''
            p = cell_val.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'

        # ---- Signatures ----
        doc.add_paragraph()
        doc.add_paragraph()
        today = date.today()

        BULAN_DICT = {
            '1': 'Januari', '2': 'Februari', '3': 'Maret', '4': 'April',
            '5': 'Mei', '6': 'Juni', '7': 'Juli', '8': 'Agustus',
            '9': 'September', '10': 'Oktober', '11': 'November', '12': 'Desember',
        }
        tgl_str = f"{today.day} {BULAN_DICT.get(str(today.month), '')} {today.year}"

        sig_table = doc.add_table(rows=4, cols=2)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_table.autofit = True

        # Remove borders from signature table
        for row_obj in sig_table.rows:
            for cell in row_obj.cells:
                for border_name in ['top', 'bottom', 'left', 'right']:
                    tag = f'w:{border_name}'
                    element = parse_xml(
                        f'<w:tcBorders {nsdecls("w")}>'
                        f'<{tag} w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                        f'</w:tcBorders>'
                    )
                    cell._tc.get_or_add_tcPr().append(element)

        # Row 0: Mengetahui / Tanggal
        for i, text in enumerate(['Mengetahui,', f'Pelaihari, {tgl_str}']):
            p = sig_table.rows[0].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'

        # Row 1: Jabatan
        for i, text in enumerate([f'Kepala {company_name}', 'Wali Kelas / Pengampu']):
            p = sig_table.rows[1].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'

        # Row 2: spacer
        for i in range(2):
            sig_table.rows[2].cells[i].text = ''
            p = sig_table.rows[2].cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(40)

        # Row 3: Names
        for i, text in enumerate(['(______________________________)', '(______________________________)']):
            p = sig_table.rows[3].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'
            run.font.underline = True

        # ---- Save & Download ----
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        docx_data = output.read()

        file_name = f"Rekap_Absensi_{self.siswa_id.name}_{self.tgl_awal}_sd_{self.tgl_akhir}.docx"

        self.write({
            'data_file': base64.b64encode(docx_data),
            'file_name': file_name
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/?model={self._name}&id={self.id}&field=data_file&download=true&filename={self.file_name}',
            'target': 'self',
        }


class WizardRekapAbsensiLine(models.TransientModel):
    _name = 'cdn.wizard_rekap_absensi_line'
    _description = 'Line Rekap Absensi Santri'

    wizard_id = fields.Many2one(
        'cdn.wizard_rekap_absensi', string='Wizard', ondelete='cascade')
    tanggal = fields.Date(string='Tanggal')
    jenis = fields.Char(string='Kegiatan')
    kehadiran = fields.Selection([
        ('Hadir', 'Hadir'),
        ('Sakit', 'Sakit'),
        ('Izin', 'Izin'),
        ('Alpa', 'Alpa'),
        ('Pulang-Sakit', 'Pulang-Sakit'),
        ('Pulang-Izin', 'Pulang-Izin'),
        ('Pulang-Alpa', 'Pulang-Alpa'),
    ], string='Kehadiran')
    pengabsen = fields.Char(string='Guru / Pengabsen')
    keterangan = fields.Char(string='Keterangan')
