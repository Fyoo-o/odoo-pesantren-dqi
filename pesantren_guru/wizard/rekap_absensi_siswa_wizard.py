from odoo import api, fields, models, _
from odoo.exceptions import UserError
import base64
import io
from datetime import date

try:
    import xlsxwriter
except ImportError:
    xlsxwriter = None

try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, Cm, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.section import WD_ORIENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class WizardRekapAbsensiSiswa(models.TransientModel):
    _name = 'cdn.wizard_rekap_absensi_siswa'
    _description = 'Wizard Rekap Absensi Siswa'
    _order = 'id desc'

    @api.model
    def default_get(self, fields_list):
        res = super(WizardRekapAbsensiSiswa, self).default_get(fields_list)
        if 'kelas_id' in fields_list and not res.get('kelas_id'):
            employee = self.env['hr.employee'].search([('user_id', '=', self.env.user.id)], limit=1)
            if employee:
                kelas = self.env['cdn.ruang_kelas'].search([('walikelas_id', '=', employee.id)], limit=1)
                if kelas:
                    res['kelas_id'] = kelas.id
                    if 'jenjang' in fields_list and hasattr(kelas, 'jenjang'):
                        res['jenjang'] = kelas.jenjang
        return res

    tgl_awal = fields.Date(string='Tanggal Awal', required=True, default=fields.Date.context_today)
    tgl_akhir = fields.Date(string='Tanggal Akhir', required=True, default=fields.Date.context_today)
    
    tipe_absensi = fields.Selection(
        selection=[
            ('kelas', 'Absensi Kelas (KBM)'),
            ('halaqoh', 'Absensi Halaqoh')
        ],
        string="Tipe Absensi",
        default='kelas',
        required=True
    )
    
    jenjang = fields.Selection(
        selection=[
            ('paud', 'PAUD'),
            ('tk', 'TK/RA'),
            ('sd', 'SD/MI'),
            ('smp', 'SMP/MTS'),
            ('sma', 'SMA/MA/SMK'),
            ('nonformal', 'Non formal'),
            ('rtq', 'Rumah Tahfidz Quran')
        ],
        string="Jenjang"
    )
    kelas_id = fields.Many2one('cdn.ruang_kelas', string='Kelas')
    halaqoh_id = fields.Many2one('cdn.halaqoh', string='Halaqoh')

    rekap_line_ids = fields.One2many(
        'cdn.wizard_rekap_absensi_siswa_line', 'wizard_id', string='Detail Kehadiran', readonly=True)

    data_file = fields.Binary(string='File')
    file_name = fields.Char(string='Nama File')

    @api.onchange('tipe_absensi', 'tgl_awal', 'tgl_akhir', 'jenjang', 'kelas_id', 'halaqoh_id')
    def _onchange_filter_rekap(self):
        if self.tipe_absensi == 'kelas':
            self.halaqoh_id = False
        elif self.tipe_absensi == 'halaqoh':
            self.jenjang = False
            self.kelas_id = False

        self._update_rekap_lines()

    @api.onchange('jenjang')
    def _onchange_jenjang(self):
        if self.jenjang and self.kelas_id and self.kelas_id.jenjang != self.jenjang:
            self.kelas_id = False

    def _update_rekap_lines(self):
        # Reset list hasil rekap
        self.rekap_line_ids = [(5, 0, 0)]

        # Validasi tanggal dan kelengkapan filter
        if not self.tgl_awal or not self.tgl_akhir or self.tgl_awal > self.tgl_akhir:
            return

        if self.tipe_absensi == 'kelas' and not self.kelas_id:
            return

        if self.tipe_absensi == 'halaqoh' and not self.halaqoh_id:
            return

        # Map: siswa_id -> dict of accumulated totals
        siswa_summary = {}

        def init_siswa_dict(siswa):
            rk = getattr(siswa, 'ruang_kelas_id', False) or getattr(siswa, 'kelas_id', False)
            return {
                'siswa_id': siswa.id,
                'siswa_name': siswa.name or '',
                'kelas_id': rk.id if rk else False,
                'jenjang': rk.jenjang if rk and hasattr(rk, 'jenjang') else False,
                'walikelas_id': rk.walikelas_id.id if rk and hasattr(rk, 'walikelas_id') and rk.walikelas_id else False,
                'hadir': 0,
                'sakit': 0,
                'izin': 0,
                'alpa': 0,
                'pulang_sakit': 0,
                'pulang_izin': 0,
                'pulang_alpa': 0,
                'keluar': 0,
            }

        # 1. Data Absensi Kelas / KBM
        if self.tipe_absensi == 'kelas':
            domain_kbm = [
                ('tanggal', '>=', self.tgl_awal),
                ('tanggal', '<=', self.tgl_akhir),
                ('kelas_id', '=', self.kelas_id.id)
            ]
            if self.jenjang:
                domain_kbm.append(('kelas_id.jenjang', '=', self.jenjang))
                
            kbm_lines = self.env['cdn.absensi_siswa_lines'].sudo().search(domain_kbm)
            for line in kbm_lines:
                if not line.siswa_id:
                    continue
                sid = line.siswa_id.id
                if sid not in siswa_summary:
                    siswa_summary[sid] = init_siswa_dict(line.siswa_id)
                
                st = line.kehadiran
                if st == 'Hadir':
                    siswa_summary[sid]['hadir'] += 1
                elif st == 'Sakit':
                    siswa_summary[sid]['sakit'] += 1
                elif st == 'Izin':
                    siswa_summary[sid]['izin'] += 1
                elif st == 'Alpa':
                    siswa_summary[sid]['alpa'] += 1
                elif st == 'Pulang-Sakit':
                    siswa_summary[sid]['pulang_sakit'] += 1
                elif st == 'Pulang-Izin':
                    siswa_summary[sid]['pulang_izin'] += 1
                elif st == 'Pulang-Alpa':
                    siswa_summary[sid]['pulang_alpa'] += 1
                elif st == 'keluar':
                    siswa_summary[sid]['keluar'] += 1

        # 2. Data Absensi Halaqoh
        elif self.tipe_absensi == 'halaqoh':
            domain_halaqoh = [
                ('tanggal', '>=', self.tgl_awal),
                ('tanggal', '<=', self.tgl_akhir),
                ('halaqoh_id', '=', self.halaqoh_id.id)
            ]
            halaqoh_lines = self.env['cdn.absen_halaqoh_line'].sudo().search(domain_halaqoh)
            for line in halaqoh_lines:
                if not line.siswa_id:
                    continue
                sid = line.siswa_id.id
                if sid not in siswa_summary:
                    siswa_summary[sid] = init_siswa_dict(line.siswa_id)
                
                st = line.kehadiran
                if st == 'Hadir':
                    siswa_summary[sid]['hadir'] += 1
                elif st == 'Sakit':
                    siswa_summary[sid]['sakit'] += 1
                elif st == 'Izin':
                    siswa_summary[sid]['izin'] += 1
                elif st == 'Alpa':
                    siswa_summary[sid]['alpa'] += 1
                elif st == 'Pulang-Sakit':
                    siswa_summary[sid]['pulang_sakit'] += 1
                elif st == 'Pulang-Izin':
                    siswa_summary[sid]['pulang_izin'] += 1
                elif st == 'Pulang-Alpa':
                    siswa_summary[sid]['pulang_alpa'] += 1
                elif st == 'keluar':
                    siswa_summary[sid]['keluar'] += 1

        # Sertakan seluruh siswa yang terdaftar di kelas / halaqoh tersebut
        if self.tipe_absensi == 'kelas' and self.kelas_id:
            s_domain = [('ruang_kelas_id', '=', self.kelas_id.id)]
            all_class_siswa = self.env['cdn.siswa'].sudo().search(s_domain)
            for s in all_class_siswa:
                if s.id not in siswa_summary:
                    siswa_summary[s.id] = init_siswa_dict(s)
        elif self.tipe_absensi == 'halaqoh' and self.halaqoh_id and self.halaqoh_id.siswa_ids:
            for s in self.halaqoh_id.siswa_ids:
                if s.id not in siswa_summary:
                    siswa_summary[s.id] = init_siswa_dict(s)

        lines_data = list(siswa_summary.values())
        lines_data.sort(key=lambda x: x['siswa_name'])

        for idx, d in enumerate(lines_data, 1):
            d['sequence'] = idx

        lines = [(0, 0, d) for d in lines_data]
        self.rekap_line_ids = lines

    def action_proses(self):
        if self.tgl_awal > self.tgl_akhir:
            raise UserError(_('Tanggal Awal tidak boleh lebih besar dari Tanggal Akhir.'))

        self._update_rekap_lines()

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'cdn.wizard_rekap_absensi_siswa',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_export_docx(self):
        if not HAS_DOCX:
            raise UserError(_("Modul 'python-docx' tidak ditemukan. Silakan hubungi administrator."))

        if self.tgl_awal > self.tgl_akhir:
            raise UserError(_('Tanggal Awal tidak boleh lebih besar dari Tanggal Akhir.'))

        if self.tipe_absensi == 'kelas' and not self.kelas_id:
            raise UserError(_('Silakan pilih Kelas terlebih dahulu.'))

        if self.tipe_absensi == 'halaqoh' and not self.halaqoh_id:
            raise UserError(_('Silakan pilih Halaqoh terlebih dahulu.'))

        # 1. Ambil data Absensi berdasarkan Tipe Absensi
        kbm_lines = self.env['cdn.absensi_siswa_lines']
        halaqoh_lines = self.env['cdn.absen_halaqoh_line']

        if self.tipe_absensi == 'kelas':
            domain_kbm = [
                ('tanggal', '>=', self.tgl_awal),
                ('tanggal', '<=', self.tgl_akhir),
                ('kelas_id', '=', self.kelas_id.id)
            ]
            if self.jenjang:
                domain_kbm.append(('kelas_id.jenjang', '=', self.jenjang))
            kbm_lines = self.env['cdn.absensi_siswa_lines'].sudo().search(domain_kbm)

        elif self.tipe_absensi == 'halaqoh':
            domain_halaqoh = [
                ('tanggal', '>=', self.tgl_awal),
                ('tanggal', '<=', self.tgl_akhir),
                ('halaqoh_id', '=', self.halaqoh_id.id)
            ]
            halaqoh_lines = self.env['cdn.absen_halaqoh_line'].sudo().search(domain_halaqoh)

        if not kbm_lines and not halaqoh_lines:
            raise UserError(_("Tidak ada data presensi pada rentang tanggal dan filter tersebut."))

        # 2. Kumpulkan daftar Tanggal Unik & daftar Santri Unik
        dates_set = set()
        siswa_set = set()
        
        attendance_map = {}
        guru_pengganti_map = {}
        keterangan_map = {}

        PRESENCE_CODE = {
            'Hadir': 'H',
            'Sakit': 'S',
            'Izin': 'I',
            'Alpa': 'A',
            'Pulang-Sakit': 'PS',
            'Pulang-Izin': 'PI',
            'Pulang-Alpa': 'PA',
            'keluar': 'K',
        }

        def clean_note(note_str):
            if not note_str or not str(note_str).strip():
                return '-'
            s = str(note_str).strip().lower()
            if s.startswith('absensi halaqoh') or s.startswith('absen halaqoh') or s.startswith('absensi kelas'):
                return '-'
            return str(note_str).strip()

        for line in kbm_lines:
            if not line.siswa_id or not line.tanggal:
                continue
            dates_set.add(line.tanggal)
            siswa_set.add(line.siswa_id)
            attendance_map[(line.siswa_id.id, line.tanggal)] = PRESENCE_CODE.get(line.kehadiran, line.kehadiran or 'H')
            
            absen_hdr = line.absensi_id
            if absen_hdr:
                is_pengganti = getattr(absen_hdr, 'is_guru_pengganti', False)
                guru = getattr(absen_hdr, 'guru_id', False)
                if is_pengganti and guru:
                    existing_gp = guru_pengganti_map.get(line.tanggal, '-')
                    if existing_gp == '-' or not existing_gp:
                        guru_pengganti_map[line.tanggal] = guru.name
                    elif guru.name not in existing_gp:
                        guru_pengganti_map[line.tanggal] += f", {guru.name}"
                elif line.tanggal not in guru_pengganti_map:
                    guru_pengganti_map[line.tanggal] = '-'
                
                ket = getattr(absen_hdr, 'keterangan', False)
                cleaned_k = clean_note(ket)
                if cleaned_k != '-':
                    existing_ket = keterangan_map.get(line.tanggal, '-')
                    if existing_ket == '-' or not existing_ket:
                        keterangan_map[line.tanggal] = cleaned_k
                    elif cleaned_k not in existing_ket:
                        keterangan_map[line.tanggal] += f" | {cleaned_k}"
                elif line.tanggal not in keterangan_map:
                    keterangan_map[line.tanggal] = '-'

        for line in halaqoh_lines:
            if not line.siswa_id or not line.tanggal:
                continue
            dates_set.add(line.tanggal)
            siswa_set.add(line.siswa_id)
            attendance_map[(line.siswa_id.id, line.tanggal)] = PRESENCE_CODE.get(line.kehadiran, line.kehadiran or 'H')

            absen_hdr = line.absen_id
            if absen_hdr and line.tanggal not in guru_pengganti_map:
                is_pengganti = getattr(absen_hdr, 'is_guru_pengganti', False)
                ustadz = getattr(absen_hdr, 'ustadz_id', False)
                pj = getattr(absen_hdr, 'penanggung_jawab_id', False)
                
                if ustadz and (is_pengganti or (pj and ustadz.id != pj.id)):
                    guru_pengganti_map[line.tanggal] = ustadz.name
                else:
                    guru_pengganti_map[line.tanggal] = '-'
                
                ket = getattr(absen_hdr, 'keterangan', False)
                keterangan_map[line.tanggal] = clean_note(ket)

        import datetime
        sorted_dates = []
        curr_d = self.tgl_awal
        while curr_d <= self.tgl_akhir:
            sorted_dates.append(curr_d)
            curr_d += datetime.timedelta(days=1)

        if self.tipe_absensi == 'halaqoh' and self.halaqoh_id:
            h_siswa = set(self.halaqoh_id.siswa_ids) if self.halaqoh_id.siswa_ids else set()
            all_siswa = h_siswa.union(siswa_set)
            sorted_siswa = sorted(list(all_siswa), key=lambda s: s.name or '')
        elif self.tipe_absensi == 'kelas' and self.kelas_id:
            s_domain = [('ruang_kelas_id', '=', self.kelas_id.id)]
            k_siswa = set(self.env['cdn.siswa'].sudo().search(s_domain))
            all_siswa = k_siswa.union(siswa_set)
            sorted_siswa = sorted(list(all_siswa), key=lambda s: s.name or '')
        else:
            sorted_siswa = sorted(list(siswa_set), key=lambda s: (s.name or ''))

        doc = DocxDocument()

        # Landscape Orientation
        for section in doc.sections:
            section.orientation = WD_ORIENT.LANDSCAPE
            new_width, new_height = section.page_height, section.page_width
            section.page_width = new_width
            section.page_height = new_height
            section.top_margin = Cm(1.5)
            section.bottom_margin = Cm(1.5)
            section.left_margin = Cm(1.5)
            section.right_margin = Cm(1.5)

        company_name = self.env.user.company_id.name or ''

        # Title
        if self.tipe_absensi == 'kelas':
            title_text = 'REKAP ABSENSI KELAS'
        elif self.tipe_absensi == 'halaqoh':
            title_text = 'REKAP ABSENSI HALAQOH'
        else:
            title_text = 'REKAP ABSENSI KELAS / HALAQOH'

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(0)
        r_title = p_title.add_run(title_text)
        r_title.bold = True
        r_title.font.size = Pt(14)
        r_title.font.name = 'Times New Roman'

        p_company = doc.add_paragraph()
        p_company.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_company.paragraph_format.space_before = Pt(0)
        p_company.paragraph_format.space_after = Pt(12)
        r_company = p_company.add_run(company_name.upper())
        r_company.bold = True
        r_company.font.size = Pt(12)
        r_company.font.name = 'Times New Roman'

        MONTHS = {
            1: 'Januari', 2: 'Februari', 3: 'Maret', 4: 'April',
            5: 'Mei', 6: 'Juni', 7: 'Juli', 8: 'Agustus',
            9: 'September', 10: 'Oktober', 11: 'November', 12: 'Desember'
        }
        if self.tgl_awal.month == self.tgl_akhir.month and self.tgl_awal.year == self.tgl_akhir.year:
            periode_str = f"{self.tgl_awal.day} - {self.tgl_akhir.day} {MONTHS[self.tgl_akhir.month]} {self.tgl_akhir.year}"
        else:
            periode_str = f"{self.tgl_awal.day} {MONTHS[self.tgl_awal.month]} - {self.tgl_akhir.day} {MONTHS[self.tgl_akhir.month]} {self.tgl_akhir.year}"

        tipe_dict = dict(self._fields['tipe_absensi'].selection)
        info_lines = [
            ('PERIODE', periode_str),
            ('TIPE ABSENSI', tipe_dict.get(self.tipe_absensi, self.tipe_absensi)),
        ]
        if self.kelas_id:
            k_name = self.kelas_id.name.name if hasattr(self.kelas_id.name, 'name') and self.kelas_id.name.name else self.kelas_id.display_name
            info_lines.append(('KELAS', k_name))
            if hasattr(self.kelas_id, 'walikelas_id') and self.kelas_id.walikelas_id:
                info_lines.append(('WALI KELAS', self.kelas_id.walikelas_id.name))
        if self.halaqoh_id:
            info_lines.append(('HALAQOH', self.halaqoh_id.name))
            if self.halaqoh_id.penanggung_jawab_id:
                info_lines.append(('USTADZ UTAMA', self.halaqoh_id.penanggung_jawab_id.name))

        for label, value in info_lines:
            p_info = doc.add_paragraph()
            p_info.paragraph_format.space_before = Pt(0)
            p_info.paragraph_format.space_after = Pt(2)
            p_info.paragraph_format.line_spacing = 1.0
            r_label = p_info.add_run(f'{label:<20}')
            r_label.bold = True
            r_label.font.size = Pt(9)
            r_label.font.name = 'Times New Roman'
            r_val = p_info.add_run(f': {value}')
            r_val.font.size = Pt(9)
            r_val.font.name = 'Times New Roman'

        doc.add_paragraph()

        # Data Table
        headers = ['No', 'Nama Santri'] + [d.strftime('%d/%m') for d in sorted_dates] + ['H', 'S', 'I', 'A', 'PS', 'PI', 'PA']
        num_cols = len(headers)
        num_rows = len(sorted_siswa) + 1 + 2

        table = doc.add_table(rows=num_rows, cols=num_cols)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        # Header Row
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(header)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)
            run.font.name = 'Times New Roman'
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="203764"/>')
            cell._tc.get_or_add_tcPr().append(shd)

        # Student Rows
        for idx, s in enumerate(sorted_siswa):
            row_idx = idx + 1
            row_cells = table.rows[row_idx].cells
            
            p = row_cells[0].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(idx + 1))
            r.font.size = Pt(8)
            r.font.name = 'Times New Roman'

            p = row_cells[1].paragraphs[0]
            r = p.add_run(s.name or '-')
            r.font.size = Pt(8)
            r.font.name = 'Times New Roman'

            counts = {'H': 0, 'S': 0, 'I': 0, 'A': 0, 'PS': 0, 'PI': 0, 'PA': 0, 'K': 0}

            col_c = 2
            for d in sorted_dates:
                st = attendance_map.get((s.id, d), '-')
                p = row_cells[col_c].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(st)
                r.font.size = Pt(8)
                r.font.name = 'Times New Roman'
                if st in counts:
                    counts[st] += 1
                col_c += 1

            for code in ['H', 'S', 'I', 'A', 'PS', 'PI', 'PA']:
                p = row_cells[col_c].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(str(counts[code]))
                r.font.size = Pt(8)
                r.font.name = 'Times New Roman'
                col_c += 1

        # Bottom Row 1: Guru Pengganti
        gp_row_idx = len(sorted_siswa) + 1
        gp_cells = table.rows[gp_row_idx].cells
        gp_cells[0].merge(gp_cells[1])
        p = gp_cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('Guru Pengganti')
        r.bold = True
        r.font.size = Pt(8)
        r.font.name = 'Times New Roman'
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EFEFEF"/>')
        gp_cells[0]._tc.get_or_add_tcPr().append(shd)

        col_c = 2
        for d in sorted_dates:
            g_name = guru_pengganti_map.get(d, '-')
            p = gp_cells[col_c].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(g_name)
            r.font.size = Pt(7)
            r.font.name = 'Times New Roman'
            col_c += 1

        # Bottom Row 2: Keterangan
        ket_row_idx = len(sorted_siswa) + 2
        ket_cells = table.rows[ket_row_idx].cells
        ket_cells[0].merge(ket_cells[1])
        p = ket_cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('Keterangan')
        r.bold = True
        r.font.size = Pt(8)
        r.font.name = 'Times New Roman'
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EFEFEF"/>')
        ket_cells[0]._tc.get_or_add_tcPr().append(shd)

        col_c = 2
        for d in sorted_dates:
            ket_val = keterangan_map.get(d, '-')
            p = ket_cells[col_c].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(ket_val)
            r.font.size = Pt(7)
            r.font.name = 'Times New Roman'
            col_c += 1

        # Legend Table
        doc.add_paragraph()
        p_leg = doc.add_paragraph()
        p_leg.paragraph_format.space_after = Pt(4)
        r_leg = p_leg.add_run('KETERANGAN KODE PRESENSI')
        r_leg.bold = True
        r_leg.font.size = Pt(10)
        r_leg.font.name = 'Times New Roman'

        legends = [
            ('H', 'Hadir'), ('S', 'Sakit'), ('I', 'Izin'), ('A', 'Alpa'),
            ('PS', 'Pulang - Sakit'), ('PI', 'Pulang - Izin'), ('PA', 'Pulang - Alpa'), ('K', 'Izin Keluar')
        ]

        leg_table = doc.add_table(rows=len(legends) + 1, cols=2)
        leg_table.style = 'Table Grid'
        leg_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        leg_table.autofit = False

        for row_obj in leg_table.rows:
            row_obj.cells[0].width = Cm(2.0)
            row_obj.cells[1].width = Cm(5.0)

        for i, header in enumerate(['KODE', 'KETERANGAN']):
            cell = leg_table.rows[0].cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(header)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)
            run.font.name = 'Times New Roman'
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="203764"/>')
            cell._tc.get_or_add_tcPr().append(shd)

        for idx, (code, name) in enumerate(legends):
            r_i = idx + 1
            c_code = leg_table.rows[r_i].cells[0]
            p = c_code.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(code)
            r.bold = True
            r.font.size = Pt(8)
            r.font.name = 'Times New Roman'

            c_name = leg_table.rows[r_i].cells[1]
            p = c_name.paragraphs[0]
            r = p.add_run(name)
            r.font.size = Pt(8)
            r.font.name = 'Times New Roman'

        # Signatures
        doc.add_paragraph()
        today = date.today()
        tgl_str = f"{today.day} {MONTHS.get(today.month, '')} {today.year}"

        sig_table = doc.add_table(rows=4, cols=2)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_table.autofit = True

        for row_obj in sig_table.rows:
            for cell in row_obj.cells:
                for border_name in ['top', 'bottom', 'left', 'right']:
                    tag = f'w:{border_name}'
                    element = parse_xml(
                        f'<w:tcBorders {nsdecls("w")}>'
                        f'<{tag} w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                        f'</w:tcBorders>'
                    )
                    cell._tc.get_or_add_tcPr().append(element)

        for i, text in enumerate(['Mengetahui,', f'Tanggal: {tgl_str}']):
            p = sig_table.rows[0].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'

        for i, text in enumerate(['Kepala Pengasuhan / Kesantrian', 'Wali Kelas / Ustadz Utama']):
            p = sig_table.rows[1].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'

        for i in range(2):
            sig_table.rows[2].cells[i].text = ''
            p = sig_table.rows[2].cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(35)

        for i, text in enumerate(['(______________________________)', '(______________________________)']):
            p = sig_table.rows[3].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'
            run.font.underline = True

        # Save & Download
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        docx_data = output.read()

        target_name = ""
        if self.tipe_absensi == 'kelas':
            tipe_str = "Kelas"
            if self.kelas_id:
                k_name = self.kelas_id.name.name if (hasattr(self.kelas_id, 'name') and hasattr(self.kelas_id.name, 'name') and self.kelas_id.name.name) else (self.kelas_id.display_name or str(self.kelas_id.name))
                target_name = f" {k_name}"
        elif self.tipe_absensi == 'halaqoh':
            tipe_str = "Halaqoh"
            if self.halaqoh_id and self.halaqoh_id.name:
                target_name = f" {self.halaqoh_id.name}"
        else:
            tipe_str = "Kelas dan Halaqoh"

        tgl_awal_str = self.tgl_awal.strftime('%d-%m-%Y') if self.tgl_awal else ''
        tgl_akhir_str = self.tgl_akhir.strftime('%d-%m-%Y') if self.tgl_akhir else ''
        file_name = f"Rekap Absensi {tipe_str}{target_name} {tgl_awal_str} sd {tgl_akhir_str}.docx"

        for char in ['/', '\\', ':', '*', '?', '"', '<', '>', '|']:
            file_name = file_name.replace(char, '-')

        self.write({
            'data_file': base64.b64encode(docx_data),
            'file_name': file_name
        })

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Export Berhasil!'),
                'message': _("File DOCX '%s' telah berhasil di-generate dan berhasil diunduh.") % self.file_name,
                'type': 'success',
                'sticky': False,
                'next': {
                    'type': 'ir.actions.act_url',
                    'url': f'/web/content/?model={self._name}&id={self.id}&field=data_file&download=true&filename={self.file_name}',
                    'target': 'self',
                }
            }
        }

    def action_export_xlsx(self):
        if not xlsxwriter:
            raise UserError(_("Modul 'xlsxwriter' tidak ditemukan. Silakan hubungi administrator."))

        if self.tgl_awal > self.tgl_akhir:
            raise UserError(_('Tanggal Awal tidak boleh lebih besar dari Tanggal Akhir.'))

        if self.tipe_absensi == 'kelas' and not self.kelas_id:
            raise UserError(_('Silakan pilih Kelas terlebih dahulu.'))

        if self.tipe_absensi == 'halaqoh' and not self.halaqoh_id:
            raise UserError(_('Silakan pilih Halaqoh terlebih dahulu.'))

        # 1. Ambil data Absensi berdasarkan Tipe Absensi
        kbm_lines = self.env['cdn.absensi_siswa_lines']
        halaqoh_lines = self.env['cdn.absen_halaqoh_line']

        if self.tipe_absensi == 'kelas':
            domain_kbm = [
                ('tanggal', '>=', self.tgl_awal),
                ('tanggal', '<=', self.tgl_akhir),
                ('kelas_id', '=', self.kelas_id.id)
            ]
            if self.jenjang:
                domain_kbm.append(('kelas_id.jenjang', '=', self.jenjang))
            kbm_lines = self.env['cdn.absensi_siswa_lines'].sudo().search(domain_kbm)

        elif self.tipe_absensi == 'halaqoh':
            domain_halaqoh = [
                ('tanggal', '>=', self.tgl_awal),
                ('tanggal', '<=', self.tgl_akhir),
                ('halaqoh_id', '=', self.halaqoh_id.id)
            ]
            halaqoh_lines = self.env['cdn.absen_halaqoh_line'].sudo().search(domain_halaqoh)

        if not kbm_lines and not halaqoh_lines:
            raise UserError(_("Tidak ada data presensi pada rentang tanggal dan filter tersebut."))

        # 2. Kumpulkan daftar Tanggal Unik & daftar Santri Unik
        dates_set = set()
        siswa_set = set()
        
        # Map: (siswa_id, tanggal) -> kehadiran_code (H, S, I, A, PS, PI, PA, K)
        attendance_map = {}
        # Map: tanggal -> guru_pengganti_name
        guru_pengganti_map = {}
        # Map: tanggal -> keterangan_sesi
        keterangan_map = {}

        PRESENCE_CODE = {
            'Hadir': 'H',
            'Sakit': 'S',
            'Izin': 'I',
            'Alpa': 'A',
            'Pulang-Sakit': 'PS',
            'Pulang-Izin': 'PI',
            'Pulang-Alpa': 'PA',
            'keluar': 'K',
        }

        def clean_note(note_str):
            if not note_str or not str(note_str).strip():
                return '-'
            s = str(note_str).strip().lower()
            if s.startswith('absensi halaqoh') or s.startswith('absen halaqoh') or s.startswith('absensi kelas'):
                return '-'
            return str(note_str).strip()

        for line in kbm_lines:
            if not line.siswa_id or not line.tanggal:
                continue
            dates_set.add(line.tanggal)
            siswa_set.add(line.siswa_id)
            attendance_map[(line.siswa_id.id, line.tanggal)] = PRESENCE_CODE.get(line.kehadiran, line.kehadiran or 'H')
            
            absen_hdr = line.absensi_id
            if absen_hdr:
                is_pengganti = getattr(absen_hdr, 'is_guru_pengganti', False)
                guru = getattr(absen_hdr, 'guru_id', False)
                if is_pengganti and guru:
                    existing_gp = guru_pengganti_map.get(line.tanggal, '-')
                    if existing_gp == '-' or not existing_gp:
                        guru_pengganti_map[line.tanggal] = guru.name
                    elif guru.name not in existing_gp:
                        guru_pengganti_map[line.tanggal] += f", {guru.name}"
                elif line.tanggal not in guru_pengganti_map:
                    guru_pengganti_map[line.tanggal] = '-'
                
                ket = getattr(absen_hdr, 'keterangan', False)
                cleaned_k = clean_note(ket)
                if cleaned_k != '-':
                    existing_ket = keterangan_map.get(line.tanggal, '-')
                    if existing_ket == '-' or not existing_ket:
                        keterangan_map[line.tanggal] = cleaned_k
                    elif cleaned_k not in existing_ket:
                        keterangan_map[line.tanggal] += f" | {cleaned_k}"
                elif line.tanggal not in keterangan_map:
                    keterangan_map[line.tanggal] = '-'

        for line in halaqoh_lines:
            if not line.siswa_id or not line.tanggal:
                continue
            dates_set.add(line.tanggal)
            siswa_set.add(line.siswa_id)
            attendance_map[(line.siswa_id.id, line.tanggal)] = PRESENCE_CODE.get(line.kehadiran, line.kehadiran or 'H')

            absen_hdr = line.absen_id
            if absen_hdr and line.tanggal not in guru_pengganti_map:
                is_pengganti = getattr(absen_hdr, 'is_guru_pengganti', False)
                ustadz = getattr(absen_hdr, 'ustadz_id', False)
                pj = getattr(absen_hdr, 'penanggung_jawab_id', False)
                
                if ustadz and (is_pengganti or (pj and ustadz.id != pj.id)):
                    guru_pengganti_map[line.tanggal] = ustadz.name
                else:
                    guru_pengganti_map[line.tanggal] = '-'
                
                ket = getattr(absen_hdr, 'keterangan', False)
                keterangan_map[line.tanggal] = clean_note(ket)

        import datetime
        sorted_dates = []
        curr_d = self.tgl_awal
        while curr_d <= self.tgl_akhir:
            sorted_dates.append(curr_d)
            curr_d += datetime.timedelta(days=1)

        if self.tipe_absensi == 'halaqoh' and self.halaqoh_id:
            h_siswa = set(self.halaqoh_id.siswa_ids) if self.halaqoh_id.siswa_ids else set()
            all_siswa = h_siswa.union(siswa_set)
            sorted_siswa = sorted(list(all_siswa), key=lambda s: s.name or '')
        elif self.tipe_absensi == 'kelas' and self.kelas_id:
            s_domain = [('ruang_kelas_id', '=', self.kelas_id.id)]
            k_siswa = set(self.env['cdn.siswa'].sudo().search(s_domain))
            all_siswa = k_siswa.union(siswa_set)
            sorted_siswa = sorted(list(all_siswa), key=lambda s: s.name or '')
        else:
            sorted_siswa = sorted(list(siswa_set), key=lambda s: (s.name or ''))

        # 3. Setup Excel Workbook
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})
        sheet = workbook.add_worksheet('Matriks Absensi')

        # Formats
        title_format = workbook.add_format({'bold': True, 'font_size': 14})
        header_format = workbook.add_format({'bold': True, 'bg_color': '#D3D3D3', 'border': 1, 'align': 'center', 'valign': 'vcenter'})
        border_format = workbook.add_format({'border': 1})
        center_format = workbook.add_format({'border': 1, 'align': 'center'})
        gray_merged_format = workbook.add_format({'bold': True, 'bg_color': '#D3D3D3', 'border': 1, 'align': 'center', 'valign': 'vcenter'})
        
        wrap_center_format = workbook.add_format({'border': 1, 'align': 'center', 'valign': 'vcenter', 'text_wrap': True, 'font_size': 9})

        # Header Info Laporan
        if self.tipe_absensi == 'kelas':
            title_text = 'REKAP ABSENSI KELAS'
        elif self.tipe_absensi == 'halaqoh':
            title_text = 'REKAP ABSENSI HALAQOH'
        else:
            title_text = 'REKAP ABSENSI KELAS / HALAQOH'
        sheet.write(0, 0, title_text, title_format)
        
        MONTHS = {
            1: 'Januari', 2: 'Februari', 3: 'Maret', 4: 'April',
            5: 'Mei', 6: 'Juni', 7: 'Juli', 8: 'Agustus',
            9: 'September', 10: 'Oktober', 11: 'November', 12: 'Desember'
        }
        if self.tgl_awal.month == self.tgl_akhir.month and self.tgl_awal.year == self.tgl_akhir.year:
            periode_str = f"Periode : {self.tgl_awal.day} - {self.tgl_akhir.day} {MONTHS[self.tgl_akhir.month]} {self.tgl_akhir.year}"
        else:
            periode_str = f"Periode : {self.tgl_awal.day} {MONTHS[self.tgl_awal.month]} - {self.tgl_akhir.day} {MONTHS[self.tgl_akhir.month]} {self.tgl_akhir.year}"

        sheet.write(2, 0, periode_str)
        tipe_dict = dict(self._fields['tipe_absensi'].selection)
        sheet.write(3, 0, f"Tipe Absensi: {tipe_dict.get(self.tipe_absensi, self.tipe_absensi)}")
        
        info_str = ""
        if self.kelas_id:
            k_name = self.kelas_id.name.name if hasattr(self.kelas_id.name, 'name') and self.kelas_id.name.name else self.kelas_id.display_name
            info_str += f"Kelas: {k_name}  "
            if hasattr(self.kelas_id, 'walikelas_id') and self.kelas_id.walikelas_id:
                info_str += f"(Wali Kelas: {self.kelas_id.walikelas_id.name})"
        if self.halaqoh_id:
            info_str += f"Halaqoh: {self.halaqoh_id.name}  "
            if self.halaqoh_id.penanggung_jawab_id:
                info_str += f"(Ustadz Utama: {self.halaqoh_id.penanggung_jawab_id.name})"
        sheet.write(4, 0, info_str)

        table_header_row = 6

        # Column widths
        sheet.set_column('A:A', 5)   # No
        sheet.set_column('B:B', 30)  # Nama Santri

        # Table Header Row
        sheet.write(table_header_row, 0, 'No', header_format)
        sheet.write(table_header_row, 1, 'Nama Santri', header_format)

        # Date Headers & Dynamic Column Widths
        col_idx = 2
        for d in sorted_dates:
            date_str = d.strftime('%d/%m')
            sheet.write(table_header_row, col_idx, date_str, header_format)
            
            g_name = guru_pengganti_map.get(d, '-')
            ket_val = keterangan_map.get(d, '-')
            
            c_width = 8
            if g_name != '-':
                c_width = max(c_width, len(str(g_name)) + 3)
            if ket_val != '-':
                c_width = max(c_width, len(str(ket_val)) + 3)
                
            sheet.set_column(col_idx, col_idx, c_width)
            col_idx += 1

        # Summary Headers
        summary_headers = ['Total H', 'Total S', 'Total I', 'Total A', 'Total PS', 'Total PI', 'Total PA']
        for sh in summary_headers:
            sheet.write(table_header_row, col_idx, sh, header_format)
            sheet.set_column(col_idx, col_idx, 10)
            col_idx += 1

        # Student Data Rows
        row_idx = table_header_row + 1
        for i, s in enumerate(sorted_siswa, 1):
            sheet.write(row_idx, 0, i, center_format)
            sheet.write(row_idx, 1, s.name or '-', border_format)

            counts = {'H': 0, 'S': 0, 'I': 0, 'A': 0, 'PS': 0, 'PI': 0, 'PA': 0, 'K': 0}

            col_c = 2
            for d in sorted_dates:
                st = attendance_map.get((s.id, d), '-')
                sheet.write(row_idx, col_c, st, center_format)
                if st in counts:
                    counts[st] += 1
                col_c += 1

            # Summary Totals
            sheet.write(row_idx, col_c, counts['H'], center_format)
            sheet.write(row_idx, col_c + 1, counts['S'], center_format)
            sheet.write(row_idx, col_c + 2, counts['I'], center_format)
            sheet.write(row_idx, col_c + 3, counts['A'], center_format)
            sheet.write(row_idx, col_c + 4, counts['PS'], center_format)
            sheet.write(row_idx, col_c + 5, counts['PI'], center_format)
            sheet.write(row_idx, col_c + 6, counts['PA'], center_format)

            row_idx += 1

        # Bottom Rows: Guru Pengganti & Keterangan
        # Bottom Row 1: Guru Pengganti
        sheet.set_row(row_idx, 28)
        sheet.merge_range(row_idx, 0, row_idx, 1, 'Guru Pengganti', gray_merged_format)
        col_c = 2
        for d in sorted_dates:
            g_name = guru_pengganti_map.get(d, '-')
            sheet.write(row_idx, col_c, g_name, wrap_center_format)
            col_c += 1
        row_idx += 1

        # Bottom Row 2: Keterangan
        sheet.set_row(row_idx, 28)
        sheet.merge_range(row_idx, 0, row_idx, 1, 'Keterangan', gray_merged_format)
        col_c = 2
        for d in sorted_dates:
            ket_val = keterangan_map.get(d, '-')
            sheet.write(row_idx, col_c, ket_val, wrap_center_format)
            col_c += 1
        row_idx += 1

        # Legenda Footer Professional (1-Column Pair Vertical Grid)
        row_idx += 2
        sheet.write(row_idx, 0, 'KETERANGAN KODE PRESENSI', workbook.add_format({'bold': True, 'font_size': 11}))
        row_idx += 1

        legend_header_format = workbook.add_format({'bold': True, 'bg_color': '#EFEFEF', 'border': 1, 'align': 'center'})
        legend_code_format = workbook.add_format({'bold': True, 'border': 1, 'align': 'center'})
        legend_text_format = workbook.add_format({'border': 1, 'align': 'left'})

        sheet.write(row_idx, 0, 'Kode', legend_header_format)
        sheet.write(row_idx, 1, 'Keterangan Presensi', legend_header_format)
        row_idx += 1

        legends = [
            ('H', 'Hadir'),
            ('S', 'Sakit'),
            ('I', 'Izin'),
            ('A', 'Alpa'),
            ('PS', 'Pulang - Sakit'),
            ('PI', 'Pulang - Izin'),
            ('PA', 'Pulang - Alpa'),
            ('K', 'Izin Keluar'),
        ]

        for code, name in legends:
            sheet.write(row_idx, 0, code, legend_code_format)
            sheet.write(row_idx, 1, name, legend_text_format)
            row_idx += 1

        workbook.close()
        output.seek(0)
        xlsx_data = output.read()

        target_name = ""
        if self.tipe_absensi == 'kelas':
            tipe_str = "Kelas"
            if self.kelas_id:
                k_name = self.kelas_id.name.name if (hasattr(self.kelas_id, 'name') and hasattr(self.kelas_id.name, 'name') and self.kelas_id.name.name) else (self.kelas_id.display_name or str(self.kelas_id.name))
                target_name = f" {k_name}"
        elif self.tipe_absensi == 'halaqoh':
            tipe_str = "Halaqoh"
            if self.halaqoh_id and self.halaqoh_id.name:
                target_name = f" {self.halaqoh_id.name}"
        else:
            tipe_str = "Kelas dan Halaqoh"
            extra = []
            if self.kelas_id:
                k_name = self.kelas_id.name.name if (hasattr(self.kelas_id, 'name') and hasattr(self.kelas_id.name, 'name') and self.kelas_id.name.name) else (self.kelas_id.display_name or str(self.kelas_id.name))
                extra.append(k_name)
            if self.halaqoh_id and self.halaqoh_id.name:
                extra.append(self.halaqoh_id.name)
            if extra:
                target_name = f" {' '.join(extra)}"

        tgl_awal_str = self.tgl_awal.strftime('%d-%m-%Y') if self.tgl_awal else ''
        tgl_akhir_str = self.tgl_akhir.strftime('%d-%m-%Y') if self.tgl_akhir else ''
        file_name = f"Rekap Absensi {tipe_str}{target_name} {tgl_awal_str} sd {tgl_akhir_str}.xlsx"

        for char in ['/', '\\', ':', '*', '?', '"', '<', '>', '|']:
            file_name = file_name.replace(char, '-')

        self.write({
            'data_file': base64.b64encode(xlsx_data),
            'file_name': file_name
        })

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Export Berhasil!'),
                'message': _("File Excel '%s' telah berhasil di-generate dan berhasil diunduh.") % self.file_name,
                'type': 'success',
                'sticky': False,
                'next': {
                    'type': 'ir.actions.act_url',
                    'url': f'/web/content/?model={self._name}&id={self.id}&field=data_file&download=true&filename={self.file_name}',
                    'target': 'self',
                }
            }
        }


class WizardRekapAbsensiSiswaLine(models.TransientModel):
    _name = 'cdn.wizard_rekap_absensi_siswa_line'
    _description = 'Detail Rekap Absensi Siswa'
    _order = 'sequence asc, tanggal asc, jenjang, kelas_id, siswa_name asc, id asc'

    wizard_id = fields.Many2one('cdn.wizard_rekap_absensi_siswa', string='Wizard', ondelete='cascade')
    sequence = fields.Integer(string='No')
    tanggal = fields.Date(string='Tanggal')
    jenjang = fields.Selection(
        selection=[
            ('paud', 'PAUD'),
            ('tk', 'TK/RA'),
            ('sd', 'SD/MI'),
            ('smp', 'SMP/MTS'),
            ('sma', 'SMA/MA/SMK'),
            ('nonformal', 'Non formal'),
            ('rtq', 'Rumah Tahfidz Quran')
        ],
        string="Jenjang"
    )
    kelas_id = fields.Many2one('cdn.ruang_kelas', string='Kelas')
    walikelas_id = fields.Many2one('hr.employee', string='Wali Kelas')
    siswa_id = fields.Many2one('cdn.siswa', string='Nama Siswa')
    siswa_name = fields.Char(string='Nama Siswa String', related='siswa_id.name', store=True)
    
    hadir = fields.Integer(string='Hadir')
    sakit = fields.Integer(string='Sakit')
    izin = fields.Integer(string='Izin')
    alpa = fields.Integer(string='Alpa')
    pulang_sakit = fields.Integer(string='Pulang Sakit')
    pulang_izin = fields.Integer(string='Pulang Izin')
    pulang_alpa = fields.Integer(string='Pulang Alpa')
    keluar = fields.Integer(string='Izin Keluar')
    guru_pengganti_name = fields.Char(string='Guru Pengganti')
