from odoo import api, fields, models, _
from odoo.exceptions import UserError
import base64
import csv
import io
from datetime import date

try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class WizardRekapAbsensiMalam(models.TransientModel):
    _name = 'cdn.wizard_rekap_absensi_malam'
    _description = 'Wizard Rekap Absensi Kamar (Musyrif)'

    def _default_musyrif_id(self):
        employee = self.env['hr.employee'].search(
            [('user_id', '=', self.env.uid)], limit=1)
        if employee and (employee.has_role('musyrif') or employee.has_role('superadmin')):
            return employee.id
        return False

    def _domain_musyrif_id(self):
        admin_user_ids = self.env.ref('base.group_system').users.ids
        return [
            '|',
            ('user_id', 'in', admin_user_ids),
            ('jns_pegawai_ids.code', 'in', ['musyrif', 'superadmin'])
        ]

    def _domain_kamar_id(self):
        tahun_ajaran = self.env.user.company_id.tahun_ajaran_aktif.id
        base_domain = [('fiscalyear_id', '=', tahun_ajaran)]

        if self.env.user.has_group('base.group_system') or self.env.user.has_group('pesantren_kesantrian.group_kesantrian_manager'):
            return base_domain

        employee = self.env['hr.employee'].search(
            [('user_id', '=', self.env.uid)], limit=1)
        if employee and (employee.has_role('musyrif') or employee.has_role('superadmin')):
            kamar_ids = self.env['cdn.kamar_santri'].search([
                ('musyrif_id', '=', employee.id),
                ('fiscalyear_id', '=', tahun_ajaran)
            ]).ids
            return [('id', 'in', kamar_ids)]
        return base_domain

    musyrif_id = fields.Many2one(
        'hr.employee', string='Musyrif/Pembina',
        domain=_domain_musyrif_id, default=_default_musyrif_id)
    kamar_id = fields.Many2one(
        'cdn.kamar_santri', string='Kamar',
        domain=_domain_kamar_id)
    tgl_awal = fields.Date(
        string='Tanggal Awal', required=True,
        default=fields.Date.context_today)
    tgl_akhir = fields.Date(
        string='Tanggal Akhir', required=True,
        default=fields.Date.context_today)

    # Result lines
    rekap_line_ids = fields.One2many(
        'cdn.wizard_rekap_absensi_malam_line', 'wizard_id',
        string='Detail Kehadiran', readonly=True)

    # Summary
    jml_hadir = fields.Integer(string='Hadir', compute='_compute_summary')
    jml_izin = fields.Integer(string='Izin', compute='_compute_summary')
    jml_sakit = fields.Integer(string='Sakit', compute='_compute_summary')
    jml_alpa = fields.Integer(string='Alpa', compute='_compute_summary')
    jml_keluar = fields.Integer(string='Izin Keluar', compute='_compute_summary')

    # Export
    data_file = fields.Binary(string='File')
    file_name = fields.Char(string='Nama File')

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        today = date.today()
        tgl_awal = res.get('tgl_awal', today)
        tgl_akhir = res.get('tgl_akhir', today)

        employee = self.env['hr.employee'].search(
            [('user_id', '=', self.env.uid)], limit=1)
        
        musyrif_id = res.get('musyrif_id')
        if not musyrif_id and employee and (employee.has_role('musyrif') or employee.has_role('superadmin')):
            musyrif_id = employee.id
            res['musyrif_id'] = musyrif_id

        kamar_id = res.get('kamar_id')
        if not kamar_id and musyrif_id:
            tahun_ajaran = self.env.user.company_id.tahun_ajaran_aktif.id
            kamar = self.env['cdn.kamar_santri'].search([
                ('musyrif_id', '=', musyrif_id),
                ('fiscalyear_id', '=', tahun_ajaran)
            ], limit=1)
            if kamar:
                kamar_id = kamar.id
                res['kamar_id'] = kamar_id

        if tgl_awal and tgl_akhir:
            lines = self._build_rekap_lines(musyrif_id, kamar_id, tgl_awal, tgl_akhir)
            res['rekap_line_ids'] = lines

        return res

    @api.depends('rekap_line_ids')
    def _compute_summary(self):
        for rec in self:
            rec.jml_hadir = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'Hadir'))
            rec.jml_izin = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'Izin'))
            rec.jml_sakit = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'Sakit'))
            rec.jml_alpa = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'Alpa'))
            rec.jml_keluar = len(rec.rekap_line_ids.filtered(
                lambda x: x.kehadiran == 'keluar'))

    @api.onchange('musyrif_id')
    def _onchange_musyrif_id(self):
        """Auto set kamar pertama dari musyrif yang dipilih dan filter domain dropdown kamar."""
        tahun_ajaran = self.env.user.company_id.tahun_ajaran_aktif.id
        if self.musyrif_id:
            kamars = self.env['cdn.kamar_santri'].search([
                ('fiscalyear_id', '=', tahun_ajaran),
                '|',
                ('musyrif_id', '=', self.musyrif_id.id),
                ('pengganti_ids', 'in', [self.musyrif_id.id])
            ])
            first_kamar = kamars[0].id if kamars else False
            self.kamar_id = first_kamar

            if self.tgl_awal and self.tgl_akhir:
                lines = self._build_rekap_lines(self.musyrif_id.id, first_kamar, self.tgl_awal, self.tgl_akhir)
                self.rekap_line_ids = [(5, 0, 0)] + lines

            return {
                'domain': {'kamar_id': [('id', 'in', kamars.ids)]}
            }
        else:
            self.kamar_id = False
            if self.tgl_awal and self.tgl_akhir:
                lines = self._build_rekap_lines(False, False, self.tgl_awal, self.tgl_akhir)
                self.rekap_line_ids = [(5, 0, 0)] + lines

            return {
                'domain': {'kamar_id': [('fiscalyear_id', '=', tahun_ajaran)]}
            }

    @api.onchange('musyrif_id', 'kamar_id', 'tgl_awal', 'tgl_akhir')
    def _onchange_rekap_params(self):
        if not (self.tgl_awal and self.tgl_akhir):
            self.rekap_line_ids = [(5, 0, 0)]
            return

        if self.tgl_awal > self.tgl_akhir:
            return

        musyrif_id = self.musyrif_id.id if self.musyrif_id else False
        kamar_id = self.kamar_id.id if self.kamar_id else False

        lines = self._build_rekap_lines(musyrif_id, kamar_id, self.tgl_awal, self.tgl_akhir)
        self.rekap_line_ids = [(5, 0, 0)] + lines

    def _build_rekap_lines(self, musyrif_id, kamar_id, tgl_awal, tgl_akhir):
        if not (kamar_id or musyrif_id):
            return []

        domain = [
            ('tanggal', '>=', tgl_awal),
            ('tanggal', '<=', tgl_akhir)
        ]

        if kamar_id:
            domain.extend(['|', ('kamar_id', '=', kamar_id), ('absen_id.kamar_id', '=', kamar_id)])
        elif musyrif_id:
            kamar_ids = self.env['cdn.kamar_santri'].search([
                ('fiscalyear_id', '=', self.env.user.company_id.tahun_ajaran_aktif.id),
                '|',
                ('musyrif_id', '=', musyrif_id),
                ('pengganti_ids', 'in', [musyrif_id])
            ]).ids
            if kamar_ids:
                domain.extend(['|', ('kamar_id', 'in', kamar_ids), ('absen_id.kamar_id', 'in', kamar_ids)])
            else:
                return []

        malam_lines = self.env['cdn.absensi_malam_line'].search(
            domain, order='tanggal desc, kamar_id asc, name asc')

        lines = []
        no = 1
        if malam_lines:
            for line in malam_lines:
                lines.append((0, 0, {
                    'no': no,
                    'tanggal': line.tanggal,
                    'kamar_id': line.kamar_id.id if line.kamar_id else (line.absen_id.kamar_id.id if line.absen_id else False),
                    'siswa_id': line.siswa_id.id if line.siswa_id else False,
                    'nis': line.nis or '',
                    'nama': line.name or '',
                    'kehadiran': line.kehadiran_absen or 'Hadir',
                    'keterangan': line.keterangan or ''
                }))
                no += 1
        else:
            # Fallback jika belum ada data absensi: daftar santri di kamar/musyrif tersebut
            santri_list = self.env['cdn.siswa']
            if kamar_id:
                kamar_rec = self.env['cdn.kamar_santri'].browse(kamar_id)
                santri_list = kamar_rec.siswa_ids
            elif musyrif_id:
                kamar_recs = self.env['cdn.kamar_santri'].search([('musyrif_id', '=', musyrif_id)])
                santri_list = kamar_recs.mapped('siswa_ids')

            for santri in santri_list.sorted(key=lambda s: s.name or ''):
                lines.append((0, 0, {
                    'no': no,
                    'tanggal': tgl_akhir,
                    'kamar_id': santri.kamar_id.id if santri.kamar_id else kamar_id,
                    'siswa_id': santri.id,
                    'nis': santri.nis or '',
                    'nama': santri.name or '',
                    'kehadiran': False,
                    'keterangan': 'Belum Diabsen'
                }))
                no += 1

        return lines

    def action_export_csv(self):
        self._onchange_rekap_params()
        if not self.rekap_line_ids:
            raise UserError(_("Tidak ada data untuk diekspor."))

        output = io.StringIO()
        writer = csv.writer(output, delimiter=',',
                            quotechar='"', quoting=csv.QUOTE_MINIMAL)

        writer.writerow(['REKAP ABSENSI KAMAR SANTRI'])
        writer.writerow([self.env.user.company_id.name or ''])
        writer.writerow([f'Musyrif: {self.musyrif_id.name or "Semua"}'])
        writer.writerow([f'Kamar: {self.kamar_id.display_name or "Semua"}'])
        writer.writerow([f'Periode: {self.tgl_awal} s/d {self.tgl_akhir}'])
        writer.writerow([])

        headers = ['No', 'Tanggal', 'Kamar', 'NIS', 'Nama Santri', 'Kehadiran', 'Keterangan']
        writer.writerow(headers)

        for line in self.rekap_line_ids:
            writer.writerow([
                line.no, line.tanggal,
                line.kamar_id.display_name if line.kamar_id else '-',
                line.nis or '-', line.nama or '-',
                dict(line._fields['kehadiran'].selection).get(line.kehadiran, line.kehadiran) if line.kehadiran else '-',
                line.keterangan or '-'
            ])

        writer.writerow([])
        writer.writerow(['Ringkasan Kehadiran'])
        writer.writerow(['Hadir', self.jml_hadir])
        writer.writerow(['Izin', self.jml_izin])
        writer.writerow(['Sakit', self.jml_sakit])
        writer.writerow(['Alpa', self.jml_alpa])
        writer.writerow(['Izin Keluar', self.jml_keluar])

        csv_data = output.getvalue().encode('utf-8')
        kamar_str = self.kamar_id.display_name if self.kamar_id else 'Semua'
        file_name = f"Rekap_Absensi_Kamar_{kamar_str}_{self.tgl_awal}_sd_{self.tgl_akhir}.csv"

        self.write({
            'data_file': base64.b64encode(csv_data),
            'file_name': file_name
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/?model={self._name}&id={self.id}&field=data_file&download=true&filename={self.file_name}',
            'target': 'self',
        }

    def action_export_docx(self):
        if not HAS_DOCX:
            raise UserError(
                _("Modul 'python-docx' tidak ditemukan. Silakan hubungi administrator."))

        self._onchange_rekap_params()
        if not self.rekap_line_ids:
            raise UserError(_("Tidak ada data untuk diekspor."))

        doc = DocxDocument()

        # Page margins
        for section in doc.sections:
            section.top_margin = Cm(1.5)
            section.bottom_margin = Cm(1.5)
            section.left_margin = Cm(1.5)
            section.right_margin = Cm(1.5)

        company_name = self.env.user.company_id.name or ''

        # ---- Title ----
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(0)
        r_title = p_title.add_run('LAPORAN REKAP ABSENSI KAMAR SANTRI')
        r_title.bold = True
        r_title.font.size = Pt(14)
        r_title.font.name = 'Times New Roman'

        p_company = doc.add_paragraph()
        p_company.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_company.paragraph_format.space_before = Pt(0)
        p_company.paragraph_format.space_after = Pt(12)
        r_company = p_company.add_run(company_name.upper())
        r_company.bold = True
        r_company.font.size = Pt(12)
        r_company.font.name = 'Times New Roman'

        # ---- Info Lines ----
        info_data = [
            ('MUSYRIF / PEMBINA', self.musyrif_id.name if self.musyrif_id else 'Semua Musyrif'),
            ('KAMAR', self.kamar_id.display_name if self.kamar_id else 'Semua Kamar'),
            ('PERIODE', f'{self.tgl_awal} s/d {self.tgl_akhir}'),
        ]
        for label, value in info_data:
            p_info = doc.add_paragraph()
            p_info.paragraph_format.space_before = Pt(0)
            p_info.paragraph_format.space_after = Pt(2)
            p_info.paragraph_format.line_spacing = 1.0
            r_label = p_info.add_run(f'{label:<24}')
            r_label.bold = True
            r_label.font.size = Pt(10)
            r_label.font.name = 'Times New Roman'
            r_val = p_info.add_run(f': {value}')
            r_val.font.size = Pt(10)
            r_val.font.name = 'Times New Roman'

        doc.add_paragraph()

        # ---- Data Table ----
        headers = ['NO', 'TANGGAL', 'KAMAR', 'NIS', 'NAMA SANTRI', 'KEHADIRAN', 'KETERANGAN']
        col_widths = [Cm(1.0), Cm(2.5), Cm(3.5), Cm(2.5), Cm(4.0), Cm(2.5), Cm(3.5)]

        num_rows = len(self.rekap_line_ids) + 1
        table = doc.add_table(rows=num_rows, cols=7)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Set column widths
        for i, width in enumerate(col_widths):
            for row_obj in table.rows:
                row_obj.cells[i].width = width

        # Header row
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(header)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'
            # Blue background
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="203764"/>')
            cell._tc.get_or_add_tcPr().append(shd)

        # Data rows
        for idx, line in enumerate(self.rekap_line_ids):
            row_idx = idx + 1
            keh_label = dict(line._fields['kehadiran'].selection).get(
                line.kehadiran, line.kehadiran) if line.kehadiran else '-'
            data = [
                str(line.no),
                str(line.tanggal) if line.tanggal else '',
                line.kamar_id.display_name if line.kamar_id else '-',
                line.nis or '-',
                line.nama or '-',
                keh_label,
                line.keterangan or '-'
            ]
            for col_idx, val in enumerate(data):
                cell = table.rows[row_idx].cells[col_idx]
                cell.text = ''
                p = cell.paragraphs[0]
                if col_idx in [0, 1, 3, 5]:  # NO, TANGGAL, NIS, KEHADIRAN centered
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(val)
                run.font.size = Pt(9)
                run.font.name = 'Times New Roman'

        # ---- Summary Table ----
        doc.add_paragraph()
        p_sum = doc.add_paragraph()
        p_sum.paragraph_format.space_after = Pt(6)
        r_sum = p_sum.add_run('RINGKASAN KEHADIRAN')
        r_sum.bold = True
        r_sum.font.size = Pt(11)
        r_sum.font.name = 'Times New Roman'

        summary_data = [
            ('Hadir', self.jml_hadir),
            ('Izin', self.jml_izin),
            ('Sakit', self.jml_sakit),
            ('Alpa', self.jml_alpa),
            ('Izin Keluar', self.jml_keluar),
        ]

        sum_table = doc.add_table(rows=len(summary_data) + 1, cols=2)
        sum_table.style = 'Table Grid'
        sum_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        sum_table.autofit = False

        # Summary column widths
        for row_obj in sum_table.rows:
            row_obj.cells[0].width = Cm(5.0)
            row_obj.cells[1].width = Cm(2.5)

        # Summary header
        for i, header in enumerate(['STATUS KEHADIRAN', 'JUMLAH']):
            cell = sum_table.rows[0].cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(header)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="203764"/>')
            cell._tc.get_or_add_tcPr().append(shd)

        # Summary data rows
        for idx, (label, val) in enumerate(summary_data):
            row_idx = idx + 1
            cell_label = sum_table.rows[row_idx].cells[0]
            cell_label.text = ''
            p = cell_label.paragraphs[0]
            run = p.add_run(label)
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'

            cell_val = sum_table.rows[row_idx].cells[1]
            cell_val.text = ''
            p = cell_val.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            run.font.name = 'Times New Roman'

        # ---- Signatures ----
        doc.add_paragraph()
        doc.add_paragraph()
        today = date.today()

        BULAN_DICT = {
            '1': 'Januari', '2': 'Februari', '3': 'Maret', '4': 'April',
            '5': 'Mei', '6': 'Juni', '7': 'Juli', '8': 'Agustus',
            '9': 'September', '10': 'Oktober', '11': 'November', '12': 'Desember',
        }
        tgl_str = f"{today.day} {BULAN_DICT.get(str(today.month), '')} {today.year}"

        sig_table = doc.add_table(rows=4, cols=2)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_table.autofit = True

        # Remove borders from signature table
        for row_obj in sig_table.rows:
            for cell in row_obj.cells:
                for border_name in ['top', 'bottom', 'left', 'right']:
                    tag = f'w:{border_name}'
                    element = parse_xml(
                        f'<w:tcBorders {nsdecls("w")}>'
                        f'<{tag} w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                        f'</w:tcBorders>'
                    )
                    cell._tc.get_or_add_tcPr().append(element)

        # Row 0: Mengetahui / Tanggal
        for i, text in enumerate(['Mengetahui,', f'Tanggal: {tgl_str}']):
            p = sig_table.rows[0].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'

        # Row 1: Jabatan
        for i, text in enumerate(['Kepala Pengasuhan / Kesantrian', 'Musyrif Pembina']):
            p = sig_table.rows[1].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'

        # Row 2: spacer
        for i in range(2):
            sig_table.rows[2].cells[i].text = ''
            p = sig_table.rows[2].cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(40)

        # Row 3: Names
        musyrif_name = self.musyrif_id.name if self.musyrif_id else '______________________________'
        for i, text in enumerate(['(______________________________)', f'({musyrif_name})']):
            p = sig_table.rows[3].cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(10)
            run.font.name = 'Times New Roman'
            run.font.underline = True

        # ---- Save & Download ----
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        docx_data = output.read()

        kamar_str = self.kamar_id.display_name if self.kamar_id else 'Semua'
        file_name = f"Rekap_Absensi_Kamar_{kamar_str}_{self.tgl_awal}_sd_{self.tgl_akhir}.docx"

        self.write({
            'data_file': base64.b64encode(docx_data),
            'file_name': file_name
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/?model={self._name}&id={self.id}&field=data_file&download=true&filename={self.file_name}',
            'target': 'self',
        }


class WizardRekapAbsensiMalamLine(models.TransientModel):
    _name = 'cdn.wizard_rekap_absensi_malam_line'
    _description = 'Line Rekap Absensi Kamar'

    wizard_id = fields.Many2one(
        'cdn.wizard_rekap_absensi_malam', string='Wizard', ondelete='cascade')
    no = fields.Integer(string='No')
    tanggal = fields.Date(string='Tanggal')
    kamar_id = fields.Many2one('cdn.kamar_santri', string='Kamar')
    siswa_id = fields.Many2one('cdn.siswa', string='Santri')
    nis = fields.Char(string='NIS')
    nama = fields.Char(string='Nama Santri')
    kehadiran = fields.Selection([
        ('Hadir', 'Hadir'),
        ('Izin', 'Izin'),
        ('keluar', 'Izin Keluar'),
        ('Sakit', 'Sakit'),
        ('Alpa', 'Alpa'),
    ], string='Kehadiran')
    keterangan = fields.Char(string='Keterangan')
